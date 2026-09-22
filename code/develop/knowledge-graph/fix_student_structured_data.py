#!/usr/bin/env python3
import csv, json, re
from pathlib import Path
from datetime import datetime, date, timedelta
from openpyxl import load_workbook

BASE = Path('/Users/xiaoji/Downloads/同步空间/openclaw/小技结构化数据')
ROOT = Path('/Users/xiaoji/Downloads/同步空间/openclaw/飞书备份/培训组/培训学员记录')
INFO_CSV = BASE / '学员信息.csv'
PROG_CSV = BASE / '学员培训进度.csv'
DICT_JSON = BASE / '_数据字典.json'

INFO_FIELDS = ['学员','教员','姓名','性别','机型','年龄','学历','入学时间','班制','理论考试日期','实操考试日期','入学测试','领队教员','对应咨询老师']
PROG_FIELDS = ['学员','教员','日期','时段','培训内容','完成情况','说明']
LABELS = {'姓名','年龄','学历','入学时间','班制','机型','领队教员','理论考试','实操考试','理论考试计划','实操考试计划','入学测试情况','对应咨询老师','工作单位','其他说明'}
TIME_SLOTS = {'上午','下午','晚上','全天'}


def norm(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)) and 30000 < float(v) < 60000:
        # Excel serial date, Windows 1900 epoch.
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).strftime('%Y-%m-%d')
    s = str(v).strip()
    # normalize full datetime strings already emitted by prior parser
    m = re.match(r'^(\d{4}-\d{1,2}-\d{1,2})\s+00:00:00$', s)
    if m:
        return fmt_ymd(m.group(1))
    return s


def fmt_ymd(s):
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', str(s).strip())
    if not m:
        return str(s).strip()
    return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'


def parse_chinese_md(s, default_year=2026):
    s = str(s).strip()
    m = re.search(r'(\d{1,2})\s*月\s*(\d{1,2})\s*(?:号|日)?', s)
    if m:
        return f'{default_year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'
    m = re.search(r'^(\d{1,2})/(\d{1,2})', s)
    if m:
        return f'{default_year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'
    m = re.search(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m:
        return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'
    return s


def clean_name_and_gender(name):
    s = norm(name)
    gender = ''
    if '男' in s:
        gender = '男'
    if '女' in s:
        gender = '女'
    # 修复“姓名后多余括号”：去除括号内性别/说明，以及孤立右括号。
    s = re.sub(r'[（(]\s*[男女]\s*[）)]', '', s)
    s = re.sub(r'[（(][^）)]*[）)]', '', s)
    s = re.sub(r'[）)]$', '', s).strip()
    return s, gender


def valid_value(v):
    s = norm(v)
    return s and s not in LABELS


def kv_from_area(ws, max_row=5, max_col=13):
    kv = {}
    rows = [[norm(ws.cell(r,c).value) for c in range(1, max_col+1)] for r in range(1, max_row+1)]
    for row in rows:
        for i, cell in enumerate(row):
            if cell in LABELS:
                vals = []
                j = i + 1
                while j < len(row) and row[j] not in LABELS:
                    if valid_value(row[j]):
                        vals.append(row[j])
                    j += 1
                if vals:
                    # 对“年龄,None,29”这种布局，取紧邻标签后的首个有效值；若标签后第一格空，则继续找。
                    kv[cell] = vals[0]
    return kv


def find_xlsx(student, teacher):
    p = ROOT / teacher / (student + '.xlsx')
    if p.exists(): return p
    p = ROOT / teacher / (Path(student).name + '.xlsx')
    if p.exists(): return p
    matches = list(ROOT.rglob(Path(student).name + '.xlsx'))
    return matches[0] if matches else None


def first_sheet_with_basic(wb):
    for ws in wb.worksheets:
        vals = [norm(c.value) for row in ws.iter_rows(min_row=1, max_row=5, max_col=13) for c in row]
        if '姓名' in vals and '机型' in vals:
            return ws
    return wb.worksheets[0]


def extract_info(student, teacher):
    xlsx = find_xlsx(student, teacher)
    row = {k:'' for k in INFO_FIELDS}
    row.update({'学员': student, '教员': teacher})
    if not xlsx:
        row['姓名'] = Path(student).name
        return row
    wb = load_workbook(xlsx, data_only=True)
    ws = first_sheet_with_basic(wb)
    kv = kv_from_area(ws)
    raw_name = kv.get('姓名') or Path(student).name
    clean, gender = clean_name_and_gender(raw_name)
    row['姓名'] = clean
    row['性别'] = gender
    row['机型'] = kv.get('机型','')
    row['年龄'] = kv.get('年龄','') if str(kv.get('年龄','')).isdigit() else ''
    row['学历'] = kv.get('学历','') if kv.get('学历','') != '学历' else ''
    row['入学时间'] = parse_chinese_md(kv.get('入学时间','')) if kv.get('入学时间','') != '入学时间' else ''
    row['班制'] = kv.get('班制','') if kv.get('班制','') != '班制' else ''
    row['理论考试日期'] = parse_chinese_md(kv.get('理论考试') or kv.get('理论考试计划',''))
    row['实操考试日期'] = parse_chinese_md(kv.get('实操考试') or kv.get('实操考试计划',''))
    row['入学测试'] = kv.get('入学测试情况','')
    row['领队教员'] = kv.get('领队教员','') if kv.get('领队教员','') != '领队教员' else ''
    row['对应咨询老师'] = kv.get('对应咨询老师','')
    return row


def is_plan_header(row):
    vals = [norm(v) for v in row]
    return ('日期' in vals[:3] and any(v in ('培训计划','培训内容') for v in vals[:5]))


def extract_progress(student, teacher):
    xlsx = find_xlsx(student, teacher)
    out = []
    if not xlsx:
        return out
    wb = load_workbook(xlsx, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        starts = [i for i, r in enumerate(rows) if is_plan_header(r)]
        for start in starts:
            last_date = ''
            for r in rows[start+1:]:
                vals = [norm(v) for v in r[:7]]
                if not any(vals):
                    # 计划区中间可能有空行；连续到下个大区时通常第一列为标题，保守跳过空行
                    continue
                if vals[0] in {'学习全流程完成情况统计','培训计划','日期','学员姓名'}:
                    continue
                # stop on obvious non-plan table title/header
                if vals[0] in {'云技学员学习全流程完成情况统计表','学员基本信息'}:
                    break
                date_s = vals[0]
                slot = vals[1]
                if date_s:
                    parsed_date = parse_chinese_md(date_s)
                    if re.match(r'^\d{4}-\d{2}-\d{2}$', parsed_date):
                        last_date = parsed_date
                    else:
                        # “第1次/在家”等不是日期，保留到说明，日期继续沿用上一条可识别日期。
                        if parsed_date and parsed_date != last_date:
                            vals[6] = ('；'.join([vals[6], parsed_date]) if vals[6] else parsed_date)
                date_out = last_date
                if slot not in TIME_SLOTS:
                    # Ignore malformed rows like copied basic-info labels.
                    continue
                content = vals[2]
                # 跳过完全空的计划内容行，解决“培训内容为空”。
                if not content:
                    continue
                complete = vals[4] if len(vals) > 4 else ''
                note_parts = []
                for idx in (5, 6):
                    if idx < len(vals) and vals[idx]:
                        note_parts.append(vals[idx])
                out.append({'学员': student, '教员': teacher, '日期': date_out, '时段': slot, '培训内容': content, '完成情况': complete, '说明': '；'.join(note_parts)})
    return out


def read_old_info():
    with open(INFO_CSV, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

old_rows = read_old_info()
new_info = [extract_info(r['学员'], r['教员']) for r in old_rows]
new_prog = []
for r in old_rows:
    new_prog.extend(extract_progress(r['学员'], r['教员']))

with open(INFO_CSV, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=INFO_FIELDS)
    w.writeheader(); w.writerows(new_info)
with open(PROG_CSV, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=PROG_FIELDS)
    w.writeheader(); w.writerows(new_prog)

with open(DICT_JSON, encoding='utf-8') as f:
    d = json.load(f)
for filename, fields, desc, rows in [
    ('学员信息.csv', INFO_FIELDS, '学员基本信息（姓名/性别/机型/年龄/学历/入学时间/考试日期等；已按原始xlsx标签-值布局修复错位）', len(new_info)),
    ('学员培训进度.csv', PROG_FIELDS, '学员每日培训记录（日期已统一为YYYY-MM-DD；空培训内容行已剔除；上午/下午继承同一日期）', len(new_prog)),
]:
    d[filename] = {'说明': desc, '行数': rows, '字段': fields, '全部字段数': len(fields)}
with open(DICT_JSON, 'w', encoding='utf-8') as f:
    json.dump(d, f, ensure_ascii=False, indent=2)

print('wrote info', len(new_info), 'progress', len(new_prog))
