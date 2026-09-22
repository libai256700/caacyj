#!/usr/bin/env python3
"""Build the four flight academy manuals from their Markdown sources."""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MANUAL_DIR = Path(__file__).resolve().parent.parent
MANUALS = (
    "学员端操作手册",
    "企业端操作手册",
    "教员端操作手册",
    "管理平台操作手册",
)

# compact_reference_guide tokens
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
HEADER_FOOTER_IN = 0.492
CONTENT_WIDTH_IN = 6.5
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BODY_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"
BODY_SIZE_PT = 11
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203748"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "C9D4E2"
WHITE = "FFFFFF"

# Named layout overrides used consistently by all four manuals.
TABLE_TEXT_SIZE_PT = 10.0
TABLE_LINE_SPACING = 1.10
MOBILE_IMAGE_MAX_HEIGHT_IN = 6.15
IMAGE_MAX_WIDTH_IN = CONTENT_WIDTH_IN
IMAGE_MAX_HEIGHT_IN = 6.15

IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
LIST_RE = re.compile(r"^(\s*)([-+*]|\d+\.)\s+(.*)$")
TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
INLINE_RE = re.compile(r"(\*\*.*?\*\*|`.*?`)")
PLACEHOLDER_RE = re.compile(r"\b(?:TODO|TBD)\b|待补充|占位符", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    east_asia: str = EAST_ASIA_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style_font(style, size: float, color: str | None = None, bold: bool | None = None) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), BODY_FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), BODY_FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), BODY_FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def add_shading(element, fill: str) -> None:
    shd = element.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        element.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_left_border(paragraph, *, color: str = BLUE, size: str = "18", space: str = "8") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    left = pbdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        pbdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)


def add_inline_runs(paragraph, text: str, *, base_size: float = BODY_SIZE_PT) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor : match.start()]), size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=base_size, bold=True)
        else:
            set_run_font(
                paragraph.add_run(token[1:-1]),
                name="Consolas",
                east_asia=EAST_ASIA_FONT,
                size=max(base_size - 0.5, 9.0),
                color=DARK_BLUE,
            )
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), size=base_size)


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(HEADER_FOOTER_IN)
    section.footer_distance = Inches(HEADER_FOOTER_IN)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    configure_style_font(normal, BODY_SIZE_PT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        configure_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    for name in ("Manual Note", "Manual Caption", "Manual TOC"):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)

    note = doc.styles["Manual Note"]
    configure_style_font(note, 10.5, INK)
    note.paragraph_format.left_indent = Inches(0.18)
    note.paragraph_format.right_indent = Inches(0.08)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(7)
    note.paragraph_format.line_spacing = 1.18
    note.paragraph_format.keep_together = True

    caption = doc.styles["Manual Caption"]
    configure_style_font(caption, 9.5, MUTED)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_together = True

    toc = doc.styles["Manual TOC"]
    configure_style_font(toc, 10.5, INK)
    toc.paragraph_format.left_indent = Inches(0.25)
    toc.paragraph_format.first_line_indent = Inches(-0.25)
    toc.paragraph_format.space_before = Pt(0)
    toc.paragraph_format.space_after = Pt(3)
    toc.paragraph_format.line_spacing = 1.08

    props = doc.core_properties
    props.title = title
    props.subject = "飞行学院前端操作手册"
    props.author = "飞行学院项目组"
    props.keywords = "飞行学院,操作手册,前端"
    props.comments = "基于 2026-08-31 项目与页面状态生成"


def add_page_field(paragraph) -> None:
    set_run_font(paragraph.add_run("第 "), size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.append(begin)
    field_run.append(instr)
    field_run.append(separate)
    field_run.append(placeholder)
    field_run.append(end)
    paragraph._p.append(field_run)
    set_run_font(paragraph.add_run(" 页"), size=9, color=MUTED)


def add_running_furniture(doc: Document, title: str) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    header_p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(header_p.add_run(title), size=9, color=MUTED)
    set_run_font(header_p.add_run("\t操作手册"), size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_field(footer_p)


def add_cover(doc: Document, title: str, role: str) -> None:
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(112)
    kicker.paragraph_format.space_after = Pt(18)
    set_run_font(kicker.add_run(f"飞行学院 · {role}"), size=11, color=BLUE, bold=True)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(10)
    set_run_font(title_p.add_run(title), size=30, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(0)
    set_run_font(subtitle.add_run("平台功能与操作流程指南"), size=15, color=DARK_BLUE)

    edition = doc.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    edition.paragraph_format.space_before = Pt(104)
    edition.paragraph_format.space_after = Pt(5)
    set_run_font(edition.add_run("基于 2026-08-31 项目与页面状态"), size=11, color=INK, bold=True)

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_after = Pt(0)
    set_run_font(version.add_run("文档版本 1.0"), size=9.5, color=MUTED, italic=True)
    version.add_run().add_break(WD_BREAK.PAGE)


def add_contents(doc: Document, chapter_titles: Iterable[str]) -> None:
    heading = doc.add_paragraph("目录", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    for chapter in chapter_titles:
        paragraph = doc.add_paragraph(style="Manual TOC")
        add_inline_runs(paragraph, chapter, base_size=10.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def next_numbering_id(root, tag: str, attr: str) -> int:
    values = [int(node.get(qn(attr))) for node in root.findall(qn(tag)) if node.get(qn(attr))]
    return max(values, default=0) + 1


def add_numbering_abstract(doc: Document, kind: str) -> int:
    root = doc.part.numbering_part.element
    abstract_id = next_numbering_id(root, "w:abstractNum", "w:abstractNumId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "271")
    ppr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        first_num.addprevious(abstract)
    return abstract_id


def add_numbering_instance(doc: Document, abstract_id: int) -> int:
    root = doc.part.numbering_part.element
    num_id = next_numbering_id(root, "w:num", "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    root.append(num)
    return num_id


def set_list_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_ref)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.widow_control = True


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def table_widths(column_count: int) -> list[int]:
    patterns = {
        1: [9360],
        2: [2700, 6660],
        3: [2100, 3060, 4200],
        4: [1500, 2300, 2780, 2780],
        5: [1200, 1800, 2100, 2100, 2160],
    }
    if column_count in patterns:
        return patterns[column_count]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblBorders", "w:tblCellMar"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)
        borders.append(border)
    tbl_pr.append(borders)

    margins = OxmlElement("w:tblCellMar")
    for edge, value in CELL_MARGINS_DXA.items():
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=column_count)
    widths = table_widths(column_count)
    for row_index, values in enumerate(normalized):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = TABLE_LINE_SPACING
            if column_count > 2 and column_index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, value, base_size=TABLE_TEXT_SIZE_PT)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
                add_shading(cell._tc.get_or_add_tcPr(), LIGHT_BLUE)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_tr_pr.append(tbl_header)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)


def add_callout(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="Manual Note")
    ppr = paragraph._p.get_or_add_pPr()
    add_shading(ppr, CALLOUT_FILL)
    set_paragraph_left_border(paragraph)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        add_inline_runs(paragraph, line, base_size=10.5)
    if any(marker in " ".join(lines) for marker in ("高风险", "极高风险", "风险提示", "重要说明")):
        for run in paragraph.runs:
            run.bold = True


def add_image(doc: Document, image_path: Path, alt_text: str, keep_with_next: bool) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Missing image: {image_path}")
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    max_height = MOBILE_IMAGE_MAX_HEIGHT_IN if aspect < 0.75 else IMAGE_MAX_HEIGHT_IN
    width = min(IMAGE_MAX_WIDTH_IN, max_height * aspect)
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3 if keep_with_next else 8)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("name", alt_text)


def next_nonblank(lines: list[str], start: int) -> str:
    for index in range(start, len(lines)):
        if lines[index].strip():
            return lines[index].strip()
    return ""


def is_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(TABLE_SEPARATOR_RE.fullmatch(cell.replace(" ", "")) for cell in cells)


def flush_plain_paragraph(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = " ".join(part.strip() for part in buffer).strip()
    if not text:
        return
    style = "Manual Caption" if re.match(r"^图\s*\d+[:：]", text) else "Normal"
    paragraph = doc.add_paragraph(style=style)
    add_inline_runs(paragraph, text, base_size=9.5 if style == "Manual Caption" else BODY_SIZE_PT)


def render_markdown_body(doc: Document, source_path: Path, abstract_ids: dict[str, int]) -> None:
    lines = source_path.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    plain_buffer: list[str] = []
    active_list_kind: str | None = None
    active_num_id: int | None = None

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            active_list_kind = None
            active_num_id = None
            index += 1
            continue

        heading_match = HEADING_RE.match(stripped)
        image_match = IMAGE_RE.match(stripped)
        list_match = LIST_RE.match(raw)

        if heading_match:
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            active_list_kind = None
            active_num_id = None
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1:
                index += 1
                continue
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            paragraph = doc.add_paragraph(style=style)
            add_inline_runs(paragraph, text, base_size={"Heading 1": 16, "Heading 2": 13, "Heading 3": 12}[style])
            index += 1
            continue

        if stripped.startswith(">"):
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            active_list_kind = None
            active_num_id = None
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(doc, quote_lines)
            continue

        if stripped.startswith("|"):
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            active_list_kind = None
            active_num_id = None
            table_lines: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = split_table_row(lines[index])
                if not is_table_separator(cells):
                    table_lines.append(cells)
                index += 1
            add_markdown_table(doc, table_lines)
            continue

        if image_match:
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            active_list_kind = None
            active_num_id = None
            alt_text, relative_path = image_match.groups()
            following = next_nonblank(lines, index + 1)
            add_image(
                doc,
                (source_path.parent / relative_path).resolve(),
                alt_text,
                keep_with_next=bool(re.match(r"^图\s*\d+[:：]", following)),
            )
            index += 1
            continue

        if list_match:
            flush_plain_paragraph(doc, plain_buffer)
            plain_buffer.clear()
            marker = list_match.group(2)
            item_text = list_match.group(3).strip()
            kind = "bullet" if marker in {"-", "+", "*"} else "decimal"
            if active_list_kind != kind or active_num_id is None:
                active_list_kind = kind
                active_num_id = add_numbering_instance(doc, abstract_ids[kind])
            paragraph = doc.add_paragraph()
            set_list_numbering(paragraph, active_num_id)
            add_inline_runs(paragraph, item_text)
            index += 1
            continue

        active_list_kind = None
        active_num_id = None
        plain_buffer.append(raw.rstrip())
        index += 1

    flush_plain_paragraph(doc, plain_buffer)


def source_chapters(source_path: Path) -> list[str]:
    chapters: list[str] = []
    for line in source_path.read_text(encoding="utf-8-sig").splitlines():
        match = re.match(r"^##\s+(.*)$", line.strip())
        if match:
            chapters.append(match.group(1).strip())
    return chapters


def source_title(source_path: Path) -> str:
    for line in source_path.read_text(encoding="utf-8-sig").splitlines():
        match = re.match(r"^#\s+(.*)$", line.strip())
        if match:
            return match.group(1).strip()
    raise ValueError(f"No title found in {source_path}")


def expected_text_fragments(source_path: Path) -> list[str]:
    fragments: list[str] = []
    for raw in source_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip().rstrip("  ")
        if not line or IMAGE_RE.match(line):
            continue
        heading = HEADING_RE.match(line)
        if heading:
            fragments.append(heading.group(2))
            continue
        if line.startswith(">"):
            line = line[1:].strip()
        list_match = LIST_RE.match(line)
        if list_match:
            line = list_match.group(3)
        if line.startswith("|"):
            cells = split_table_row(line)
            if is_table_separator(cells):
                continue
            fragments.extend(cells)
            continue
        fragments.append(line)
    return [INLINE_RE.sub(lambda m: m.group(0).strip("*`"), item) for item in fragments if item]


def document_text(doc: Document) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


def audit_document(doc: Document, source_path: Path, output_path: Path) -> dict[str, int]:
    text = document_text(doc)
    compact_text = re.sub(r"\s+", " ", text)
    missing: list[str] = []
    for fragment in expected_text_fragments(source_path):
        normalized = re.sub(r"\s+", " ", fragment).strip()
        if normalized and normalized not in compact_text:
            missing.append(normalized)
    if missing:
        preview = " | ".join(missing[:8])
        raise ValueError(f"Source coverage failed for {source_path.name}: {preview}")
    if PLACEHOLDER_RE.search(text):
        raise ValueError(f"Placeholder text found in {output_path.name}")
    if PHONE_RE.search(text):
        raise ValueError(f"Potential unredacted phone number found in {output_path.name}")

    source = source_path.read_text(encoding="utf-8-sig")
    expected_images = len(re.findall(r"^!\[", source, flags=re.MULTILINE))
    expected_tables = len(re.findall(r"^\|\s*[-:]", source, flags=re.MULTILINE))
    if len(doc.inline_shapes) != expected_images:
        raise ValueError(f"Image count mismatch in {output_path.name}")
    if len(doc.tables) != expected_tables:
        raise ValueError(f"Table count mismatch in {output_path.name}: {len(doc.tables)} != {expected_tables}")
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
    }


def build_manual(stem: str) -> tuple[Path, dict[str, int]]:
    source_path = MANUAL_DIR / f"{stem}.md"
    output_path = MANUAL_DIR / f"{stem}.docx"
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    source = source_path.read_text(encoding="utf-8-sig")
    if PLACEHOLDER_RE.search(source):
        raise ValueError(f"Placeholder text found in {source_path.name}")
    if PHONE_RE.search(source):
        raise ValueError(f"Potential unredacted phone number found in {source_path.name}")

    title = source_title(source_path)
    role = stem.replace("操作手册", "")
    doc = Document()
    configure_document(doc, title)
    add_running_furniture(doc, title)
    add_cover(doc, title, role)
    add_contents(doc, source_chapters(source_path))
    abstract_ids = {
        "bullet": add_numbering_abstract(doc, "bullet"),
        "decimal": add_numbering_abstract(doc, "decimal"),
    }
    render_markdown_body(doc, source_path, abstract_ids)
    doc.save(output_path)

    reopened = Document(output_path)
    stats = audit_document(reopened, source_path, output_path)
    return output_path, stats


def main() -> int:
    results: list[tuple[Path, dict[str, int]]] = []
    for stem in MANUALS:
        results.append(build_manual(stem))
    for output_path, stats in results:
        print(
            f"BUILT\t{output_path.name}\t{output_path.stat().st_size} bytes\t"
            f"{stats['paragraphs']} paragraphs\t{stats['tables']} tables\t{stats['images']} images"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
