from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image


OUTPUT_DIR = Path(__file__).resolve().parent
SCREENSHOT_DIR = OUTPUT_DIR / "03-截图"
TARGET_DOCX = Path.home() / "Desktop" / "云技软件截图.docx"
MAX_IMAGE_WIDTH_CM = 14.63
TABLE_IMAGE_MAX_HEIGHT_CM = 17.0
BODY_IMAGE_MAX_HEIGHT_CM = 19.0
TABLE_WIDTH_CM = 15.03


@dataclass(frozen=True)
class ScreenshotSpec:
    filename: str
    heading: str
    description: str


ACCESS_AND_AGREEMENT = (
    ScreenshotSpec(
        "01-login.png",
        "1.1 登录",
        "登录页面用于学员或企业用户通过手机号和验证码进入云技科技飞行学院。",
    ),
    ScreenshotSpec(
        "02-agreement-user.png",
        "1.2 用户服务协议",
        "用户服务协议页面展示账号使用、课程训练、安全规范及知识产权等平台规则。",
    ),
    ScreenshotSpec(
        "04-privacy.png",
        "1.3 隐私协议",
        "隐私协议页面说明平台对个人信息的收集、使用、保存和保护规则。",
    ),
)

HOME = ScreenshotSpec(
    "03-home-before-login.png",
    "2. 首页",
    "首页集中展示考题测试、常用功能和 AI 助手等主要功能入口。",
)

ENTERPRISE_AND_ACCOUNT = (
    ScreenshotSpec(
        "05-enterprise-register.png",
        "3.1 企业注册",
        "企业注册页面用于填写企业基础信息、联系人信息并提交营业执照审核。",
    ),
    ScreenshotSpec(
        "06-phone-bind.png",
        "3.2 绑定手机号",
        "绑定手机号页面用于完成第三方授权账号与手机号的登录身份关联。",
    ),
)

GENERATION_ORDER = ACCESS_AND_AGREEMENT + (HOME,) + ENTERPRISE_AND_ACCOUNT


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_inputs() -> dict[str, Path]:
    expected = sorted(
        (SCREENSHOT_DIR / spec.filename for spec in GENERATION_ORDER),
        key=lambda path: path.name,
    )
    missing = [path for path in expected if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing screenshots: {missing}")

    actual = sorted(SCREENSHOT_DIR.glob("*.png"), key=lambda path: path.name)
    if actual != expected:
        raise ValueError(f"Expected exactly six ordered PNG files, got: {actual}")
    return {path.name: path for path in expected}


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")

    properties = document.core_properties
    properties.title = "云技软件截图"
    properties.subject = "云技科技飞行学院软件界面截图"
    properties.author = "云技科技"


def add_title(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run("云技软件截图")
    set_run_font(run, 22, bold=True)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), 18, bold=True)


def fitted_width_cm(image_path: Path, max_height_cm: float) -> float:
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    if pixel_width <= 0 or pixel_height <= 0:
        raise ValueError(f"Invalid image dimensions: {image_path}")
    return min(MAX_IMAGE_WIDTH_CM, max_height_cm * pixel_width / pixel_height)


def add_image_paragraph(
    container,
    image_path: Path,
    heading: str,
    max_height_cm: float,
):
    image_paragraph = container.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(6)
    image_paragraph.paragraph_format.keep_together = True
    image_run = image_paragraph.add_run()
    shape = image_run.add_picture(
        str(image_path),
        width=Cm(fitted_width_cm(image_path, max_height_cm)),
    )
    shape._inline.docPr.set("name", heading)
    shape._inline.docPr.set("descr", f"{heading} - {image_path.name}")
    return shape


def prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def add_table_group(
    document: Document,
    group_heading: str,
    specs: tuple[ScreenshotSpec, ...],
    paths: dict[str, Path],
) -> list:
    add_section_heading(document, group_heading)
    table = document.add_table(rows=0, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(TABLE_WIDTH_CM)

    shapes = []
    for spec in specs:
        row = table.add_row()
        prevent_row_split(row)
        cell = row.cells[0]
        cell.width = Cm(TABLE_WIDTH_CM)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        heading_paragraph = cell.paragraphs[0]
        heading_paragraph.paragraph_format.space_after = Pt(4)
        heading_paragraph.paragraph_format.keep_with_next = True
        set_run_font(heading_paragraph.add_run(spec.heading), 16, bold=True)

        description_paragraph = cell.add_paragraph()
        description_paragraph.paragraph_format.space_after = Pt(6)
        description_paragraph.paragraph_format.keep_with_next = True
        set_run_font(description_paragraph.add_run(spec.description), 11)

        shapes.append(
            add_image_paragraph(
                cell,
                paths[spec.filename],
                spec.heading,
                TABLE_IMAGE_MAX_HEIGHT_CM,
            )
        )

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return shapes


def add_standalone_screenshot(
    document: Document,
    spec: ScreenshotSpec,
    image_path: Path,
):
    add_section_heading(document, spec.heading)

    description_paragraph = document.add_paragraph()
    description_paragraph.paragraph_format.space_after = Pt(8)
    description_paragraph.paragraph_format.keep_with_next = True
    set_run_font(description_paragraph.add_run(spec.description), 11)

    shape = add_image_paragraph(
        document,
        image_path,
        spec.heading,
        BODY_IMAGE_MAX_HEIGHT_CM,
    )
    return shape


def main() -> None:
    paths = validate_inputs()
    document = Document()
    configure_document(document)
    add_title(document)

    shapes = []
    shapes.extend(
        add_table_group(
            document,
            "1. 访问与协议",
            ACCESS_AND_AGREEMENT,
            paths,
        )
    )
    shapes.append(add_standalone_screenshot(document, HOME, paths[HOME.filename]))
    shapes.extend(
        add_table_group(
            document,
            "3. 企业与账号",
            ENTERPRISE_AND_ACCOUNT,
            paths,
        )
    )

    TARGET_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET_DOCX)

    print(f"OUTPUT={TARGET_DOCX}")
    print(f"BYTES={TARGET_DOCX.stat().st_size}")
    print(f"SHA256={sha256(TARGET_DOCX)}")
    for index, (spec, shape) in enumerate(zip(GENERATION_ORDER, shapes, strict=True), start=1):
        path = paths[spec.filename]
        print(
            f"MEDIA_{index}={path.name}|{shape.width / 360000:.3f}x"
            f"{shape.height / 360000:.3f}cm|{sha256(path)}"
        )


if __name__ == "__main__":
    main()
