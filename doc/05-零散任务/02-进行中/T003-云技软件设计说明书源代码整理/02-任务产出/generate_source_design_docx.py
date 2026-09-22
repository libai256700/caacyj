from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from PIL import Image


OLD_BRAND = "\u6765\u7ed8"
NEW_BRAND = "\u4e91\u6280"
MEDIA_PATH = "word/media/image1.png"
HEADER_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"
)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
BUSINESS_ROOTS = (
    "code/develop/yunjikeji",
    "code/develop/yunjikeji-admin-ui",
    "code/develop/yunjikeji-server",
    "code/develop/yunjikeji-admin-server",
    "code/develop/fly-llm",
)
PRESERVED_XML_PARTS = (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
)
CREDENTIAL_PATTERNS = (
    (
        "private-key-header",
        re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----"),
    ),
    (
        "aws-access-key",
        re.compile(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b", re.IGNORECASE),
    ),
    (
        "long-sk-key",
        re.compile(r"\bsk-[A-Za-z0-9_-]{16,}\b", re.IGNORECASE),
    ),
    (
        "literal-sensitive-assignment",
        re.compile(
            r"(?:password|passwd|pwd|secret|token|api[_-]?key|apikey|"
            r"access[_-]?key|secret[_-]?key|private[_-]?key)\s*[:=]\s*"
            r"[\"'][^\"'\r\n]{4,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "literal-authorization",
        re.compile(
            r"authorization\s*[:=]\s*[\"'](?:bearer|basic)\s+"
            r"[A-Za-z0-9._+/=-]{8,}[\"']",
            re.IGNORECASE,
        ),
    ),
)
KEYWORD_PATTERN = re.compile(
    r"password|passwd|secret|token|private[ _-]?key|access[ _-]?key|"
    r"api[ _-]?key|authorization|bearer",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class SourceItem:
    order: int
    project: str
    relative_path: str
    language: str
    expected_lines: int
    expected_chars: int
    expected_sha256: str
    source_path: Path
    source_text: str

    @property
    def normalized_source_text(self) -> str:
        return normalize_newlines(self.source_text)

    @property
    def container_text(self) -> str:
        metadata = (
            f"\u5de5\u7a0b\uff1a{self.project}\n"
            f"\u8def\u5f84\uff1a{self.relative_path}\n"
            f"\u8bed\u8a00\uff1a{self.language}\n"
            "\u6e90\u7801\uff1a\n"
        )
        return metadata + self.normalized_source_text


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def normalize_newlines(text: str) -> str:
    return re.sub(r"\r\n|\r|\n", "\n", text)


def utf16_character_count(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def logical_line_count(text: str) -> int:
    return len(text.splitlines())


def strip_code_ticks(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value.startswith("`") and value.endswith("`"):
        return value[1:-1]
    return value


def discover_repo_root(script_path: Path) -> Path:
    for candidate in (script_path.parent, *script_path.parents):
        if (candidate / "code" / "develop").is_dir():
            return candidate.resolve()
    raise RuntimeError("Cannot locate repository root from script path")


def is_within(path: Path, root: Path) -> bool:
    try:
        return os.path.commonpath((str(path), str(root))) == str(root)
    except ValueError:
        return False


def load_inventory(inventory_path: Path, repo_root: Path) -> list[SourceItem]:
    inventory_text = inventory_path.read_text(encoding="utf-8")
    rows: list[SourceItem] = []

    for line in inventory_text.splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 9 or not cells[0].isdigit():
            continue

        order = int(cells[0])
        project = strip_code_ticks(cells[1])
        relative_path = strip_code_ticks(cells[2]).replace("\\", "/")
        language = strip_code_ticks(cells[3])
        expected_lines = int(cells[5].replace(",", ""))
        expected_chars = int(cells[6].replace(",", ""))
        expected_sha256 = strip_code_ticks(cells[7]).lower()

        if not relative_path.startswith("code/develop/"):
            continue
        if relative_path != f"code/develop/{project}" and not relative_path.startswith(
            f"code/develop/{project}/"
        ):
            raise RuntimeError(
                f"Inventory project/path mismatch at row {order}: {relative_path}"
            )

        source_path = (repo_root / Path(relative_path)).resolve()
        if not is_within(source_path, repo_root):
            raise RuntimeError(f"Inventory path escapes repository: {relative_path}")
        if not source_path.is_file():
            raise FileNotFoundError(source_path)

        source_bytes = source_path.read_bytes()
        source_text = source_bytes.decode("utf-8-sig", errors="strict")
        actual_sha256 = sha256_bytes(source_bytes)
        actual_lines = logical_line_count(source_text)
        actual_chars = utf16_character_count(source_text)

        if actual_sha256 != expected_sha256:
            raise RuntimeError(f"SHA-256 drift: {relative_path}")
        if actual_lines != expected_lines:
            raise RuntimeError(
                f"Line-count drift: {relative_path}: {actual_lines} != {expected_lines}"
            )
        if actual_chars != expected_chars:
            raise RuntimeError(
                f"Character-count drift: {relative_path}: "
                f"{actual_chars} != {expected_chars}"
            )
        validate_xml_characters(source_text, relative_path)

        rows.append(
            SourceItem(
                order=order,
                project=project,
                relative_path=relative_path,
                language=language,
                expected_lines=expected_lines,
                expected_chars=expected_chars,
                expected_sha256=expected_sha256,
                source_path=source_path,
                source_text=source_text,
            )
        )

    rows.sort(key=lambda item: item.order)
    if len(rows) != 18:
        raise RuntimeError(f"Expected 18 inventory rows, found {len(rows)}")
    if [item.order for item in rows] != list(range(1, 19)):
        raise RuntimeError("Inventory order must be exactly 1 through 18")
    if len({item.relative_path for item in rows}) != len(rows):
        raise RuntimeError("Inventory contains duplicate source paths")
    if {item.project for item in rows} != {
        "yunjikeji",
        "yunjikeji-admin-ui",
        "yunjikeji-server",
        "yunjikeji-admin-server",
        "fly-llm",
    }:
        raise RuntimeError("Inventory does not cover the five required projects")
    if sum(item.expected_lines for item in rows) != 8_329:
        raise RuntimeError("Inventory total line count is not 8,329")
    if sum(item.expected_chars for item in rows) != 310_972:
        raise RuntimeError("Inventory total character count is not 310,972")
    return rows


def validate_xml_characters(text: str, label: str) -> None:
    for index, char in enumerate(text):
        codepoint = ord(char)
        valid = (
            codepoint in (0x09, 0x0A, 0x0D)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        )
        if not valid:
            raise RuntimeError(
                f"Invalid XML character U+{codepoint:04X} in {label} at {index}"
            )


def scan_credentials(items: list[SourceItem]) -> tuple[list[dict], list[dict]]:
    credential_hits: list[dict] = []
    keyword_contexts: list[dict] = []
    for item in items:
        for line_number, line in enumerate(item.source_text.splitlines(), start=1):
            for pattern_name, pattern in CREDENTIAL_PATTERNS:
                if pattern.search(line):
                    credential_hits.append(
                        {
                            "path": item.relative_path,
                            "line": line_number,
                            "pattern": pattern_name,
                        }
                    )
            if KEYWORD_PATTERN.search(line):
                keyword_contexts.append(
                    {
                        "path": item.relative_path,
                        "line": line_number,
                        "text": line.strip()[:240],
                    }
                )
    return credential_hits, keyword_contexts


def paragraph_ppr_fingerprints(document: Document) -> list[bytes]:
    fingerprints: list[bytes] = []
    for paragraph in document.paragraphs:
        ppr = paragraph._p.pPr
        fingerprints.append(
            b"" if ppr is None else etree.tostring(ppr, method="c14n")
        )
    return fingerprints


def clear_paragraph_except_properties(paragraph) -> None:
    paragraph_element = paragraph._p
    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)


def write_source_container(paragraph, item: SourceItem) -> None:
    clear_paragraph_except_properties(paragraph)
    run = paragraph.add_run(item.container_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run_properties = run._element.get_or_add_rPr()
    run_fonts = run_properties.get_or_add_rFonts()
    run_fonts.set(qn("w:ascii"), "Times New Roman")
    run_fonts.set(qn("w:hAnsi"), "Times New Roman")
    run_fonts.set(qn("w:eastAsia"), "\u7b49\u7ebf")


def replace_equal_length_text_nodes(root, old: str, new: str) -> int:
    if len(old) != len(new):
        raise ValueError("Cross-node replacement requires equal-length strings")
    text_nodes = root.xpath(".//w:t")
    original_node_texts = [node.text or "" for node in text_nodes]
    combined = "".join(original_node_texts)
    replacement_count = combined.count(old)
    if replacement_count == 0:
        return 0
    replaced = combined.replace(old, new)
    offset = 0
    for node, original_text in zip(text_nodes, original_node_texts):
        node.text = replaced[offset : offset + len(original_text)]
        offset += len(original_text)
    return replacement_count


def replace_header_brand(document: Document) -> int:
    replacement_count = 0
    seen_parts: set[str] = set()
    for part in document.part.package.parts:
        if part.content_type != HEADER_CONTENT_TYPE:
            continue
        part_name = str(part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        replacement_count += replace_equal_length_text_nodes(
            part.element, OLD_BRAND, NEW_BRAND
        )
    if replacement_count < 1:
        raise RuntimeError(f"Header brand {OLD_BRAND!r} was not found")
    return replacement_count


def create_white_png() -> bytes:
    image = Image.new("RGB", (640, 960), (255, 255, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def rebuild_media(source_docx: Path, rebuilt_docx: Path, media: bytes) -> None:
    found_media = False
    with zipfile.ZipFile(source_docx, "r") as source_zip:
        with zipfile.ZipFile(rebuilt_docx, "w") as target_zip:
            target_zip.comment = source_zip.comment
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == MEDIA_PATH:
                    data = media
                    found_media = True
                target_zip.writestr(info, data)
    if not found_media:
        raise RuntimeError(f"Template package does not contain {MEDIA_PATH}")


def canonical_xml(data: bytes) -> bytes:
    return etree.tostring(etree.fromstring(data), method="c14n")


def package_validation(template_path: Path, output_path: Path) -> dict:
    with zipfile.ZipFile(template_path, "r") as template_zip:
        template_parts = {
            name for name in template_zip.namelist() if not name.endswith("/")
        }
        template_section = etree.fromstring(
            template_zip.read("word/document.xml")
        ).xpath("/w:document/w:body/w:sectPr", namespaces=NS)
        if len(template_section) != 1:
            raise RuntimeError("Template must contain one section property element")
        template_section_xml = etree.tostring(template_section[0], method="c14n")
        preserved_part_xml = {
            name: canonical_xml(template_zip.read(name)) for name in PRESERVED_XML_PARTS
        }

    with zipfile.ZipFile(output_path, "r") as output_zip:
        if output_zip.testzip() is not None:
            raise RuntimeError("Output DOCX ZIP CRC check failed")
        output_parts = {
            name for name in output_zip.namelist() if not name.endswith("/")
        }
        required_parts = {
            "[Content_Types].xml",
            "word/document.xml",
            "word/styles.xml",
            "word/_rels/document.xml.rels",
            MEDIA_PATH,
        }
        missing_parts = required_parts - output_parts
        if missing_parts:
            raise RuntimeError(f"Output package missing parts: {sorted(missing_parts)}")
        if template_parts - output_parts:
            raise RuntimeError("Output package dropped template parts")

        xml_parts = [
            name
            for name in output_zip.namelist()
            if name.endswith(".xml") or name.endswith(".rels")
        ]
        parsed_xml: dict[str, etree._Element] = {}
        for name in xml_parts:
            parsed_xml[name] = etree.fromstring(output_zip.read(name))

        document_root = parsed_xml["word/document.xml"]
        body_paragraphs = document_root.xpath(
            "/w:document/w:body/w:p", namespaces=NS
        )
        table_count = len(document_root.xpath(".//w:tbl", namespaces=NS))
        section_nodes = document_root.xpath(".//w:sectPr", namespaces=NS)
        if len(body_paragraphs) != 20 or table_count != 0 or len(section_nodes) != 1:
            raise RuntimeError(
                "Output XML structure is not 20 paragraphs / 0 tables / 1 section"
            )
        if etree.tostring(section_nodes[0], method="c14n") != template_section_xml:
            raise RuntimeError("Section properties changed from the template")

        for name, expected_xml in preserved_part_xml.items():
            if canonical_xml(output_zip.read(name)) != expected_xml:
                raise RuntimeError(f"Template style part changed: {name}")

        all_xml_text = "\n".join(
            "".join(root.itertext()) for root in parsed_xml.values()
        )
        header_names = sorted(
            name
            for name in parsed_xml
            if name.startswith("word/header") and name.endswith(".xml")
        )
        header_text = "\n".join(
            "".join(parsed_xml[name].itertext()) for name in header_names
        )
        if OLD_BRAND in all_xml_text:
            raise RuntimeError(f"Old brand remains in output XML: {OLD_BRAND}")
        if NEW_BRAND not in header_text:
            raise RuntimeError(f"New brand is missing from output headers: {NEW_BRAND}")

        media_bytes = output_zip.read(MEDIA_PATH)
        with Image.open(io.BytesIO(media_bytes)) as media_image:
            media_image.load()
            if media_image.size != (640, 960):
                raise RuntimeError("Replacement media is not 640x960")
            rgb_image = media_image.convert("RGB")
            extrema = rgb_image.getextrema()
            if extrema != ((255, 255), (255, 255), (255, 255)):
                raise RuntimeError("Replacement media is not pure white")

    return {
        "zip_crc": "ok",
        "zip_parts": len(output_parts),
        "xml_parts_parsed": len(xml_parts),
        "paragraphs": len(body_paragraphs),
        "tables": table_count,
        "sections": len(section_nodes),
        "headers": header_names,
        "old_brand_occurrences": 0,
        "new_brand_in_header": True,
        "media_path": MEDIA_PATH,
        "media_size": [640, 960],
        "media_pure_white": True,
        "style_parts_preserved": list(PRESERVED_XML_PARTS),
        "section_properties_preserved": True,
    }


def document_validation(
    template_path: Path,
    output_path: Path,
    items: list[SourceItem],
    expected_ppr: list[bytes],
) -> dict:
    template_document = Document(template_path)
    output_document = Document(output_path)
    if len(output_document.paragraphs) != 20:
        raise RuntimeError("Reopened output does not contain 20 body paragraphs")
    if len(output_document.tables) != 0 or len(output_document.sections) != 1:
        raise RuntimeError("Reopened output table/section count is incorrect")
    if output_document.paragraphs[0].text != template_document.paragraphs[0].text:
        raise RuntimeError("Opening marker changed")
    if output_document.paragraphs[-1].text != template_document.paragraphs[-1].text:
        raise RuntimeError("Closing marker changed")
    if paragraph_ppr_fingerprints(output_document) != expected_ppr:
        raise RuntimeError("One or more paragraph properties changed")

    for item, paragraph in zip(items, output_document.paragraphs[1:19]):
        if paragraph.text != item.container_text:
            raise RuntimeError(
                f"Reopened paragraph content mismatch: {item.relative_path}"
            )

    full_text = "\n".join(paragraph.text for paragraph in output_document.paragraphs)
    generated_credential_hits = []
    for pattern_name, pattern in CREDENTIAL_PATTERNS:
        if pattern.search(full_text):
            generated_credential_hits.append(pattern_name)
    if generated_credential_hits:
        raise RuntimeError(
            f"Credential patterns found after generation: {generated_credential_hits}"
        )

    return {
        "reopen": "ok",
        "paragraphs": len(output_document.paragraphs),
        "tables": len(output_document.tables),
        "sections": len(output_document.sections),
        "source_containers": 18,
        "paths_verified": len(items),
        "full_sources_verified": len(items),
        "paragraph_properties_preserved": True,
        "opening_marker": output_document.paragraphs[0].text,
        "closing_marker": output_document.paragraphs[-1].text,
        "generated_credential_hits": generated_credential_hits,
    }


def business_git_status(repo_root: Path) -> str:
    command = [
        "git",
        "status",
        "--porcelain=v1",
        "--untracked-files=all",
        "--",
        *BUSINESS_ROOTS,
    ]
    result = subprocess.run(
        command,
        cwd=repo_root,
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return result.stdout.decode("utf-8", errors="strict").strip()


def generate(
    template_path: Path,
    inventory_path: Path,
    output_path: Path,
    repo_root: Path,
) -> dict:
    if not template_path.is_file():
        raise FileNotFoundError(template_path)
    if not inventory_path.is_file():
        raise FileNotFoundError(inventory_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    template_sha256_before = sha256_file(template_path)
    items = load_inventory(inventory_path, repo_root)
    credential_hits, keyword_contexts = scan_credentials(items)
    if credential_hits:
        raise RuntimeError(f"Credential-shaped values found: {credential_hits}")

    git_before = business_git_status(repo_root)
    if git_before:
        raise RuntimeError(f"Business source worktree is not clean:\n{git_before}")

    document = Document(template_path)
    if len(document.paragraphs) != 20:
        raise RuntimeError("Template must contain exactly 20 body paragraphs")
    if len(document.tables) != 0 or len(document.sections) != 1:
        raise RuntimeError("Template must contain 0 tables and 1 section")
    expected_ppr = paragraph_ppr_fingerprints(document)
    header_replacements = replace_header_brand(document)

    for item, paragraph in zip(items, document.paragraphs[1:19]):
        write_source_container(paragraph, item)

    unique_id = uuid.uuid4().hex
    python_docx_path = output_path.parent / (
        f".{output_path.stem}.{unique_id}.python-docx.docx"
    )
    candidate_path = output_path.parent / (
        f".{output_path.stem}.{unique_id}.candidate.docx"
    )
    try:
        document.save(python_docx_path)
        rebuild_media(python_docx_path, candidate_path, create_white_png())

        package_result = package_validation(template_path, candidate_path)
        document_result = document_validation(
            template_path, candidate_path, items, expected_ppr
        )

        if sha256_file(template_path) != template_sha256_before:
            raise RuntimeError("Template changed during generation")
        for item in items:
            if sha256_file(item.source_path) != item.expected_sha256:
                raise RuntimeError(
                    f"Business source changed during generation: {item.relative_path}"
                )

        git_after = business_git_status(repo_root)
        if git_after:
            raise RuntimeError(f"Business source worktree changed:\n{git_after}")

        candidate_hash = sha256_file(candidate_path)
        candidate_size = candidate_path.stat().st_size
        os.replace(candidate_path, output_path)
        if sha256_file(output_path) != candidate_hash:
            raise RuntimeError("Final output differs from validated candidate")
        Document(output_path)

        return {
            "status": "passed",
            "template": str(template_path),
            "template_size": template_path.stat().st_size,
            "template_sha256": template_sha256_before,
            "inventory": str(inventory_path),
            "output": str(output_path),
            "output_size": candidate_size,
            "output_sha256": candidate_hash,
            "source_files": len(items),
            "source_lines": sum(item.expected_lines for item in items),
            "source_characters": sum(item.expected_chars for item in items),
            "header_brand_replacements": header_replacements,
            "credential_hits": credential_hits,
            "keyword_contexts": keyword_contexts,
            "git_business_status_before": git_before,
            "git_business_status_after": git_after,
            "package": package_result,
            "document": document_result,
        }
    finally:
        python_docx_path.unlink(missing_ok=True)
        candidate_path.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    script_path = Path(__file__).resolve()
    repo_root = discover_repo_root(script_path)
    output_dir = script_path.parent
    default_template = (
        Path(r"C:\Users\18163\Downloads")
        / "\u98de\u4e6620260724-105710"
        / "\u6765\u7ed8_\u8f6f\u4ef6\u8bbe\u8ba1\u8bf4\u660e\u4e66\u6e90\u4ee3\u7801_1.docx"
    )
    default_output = (
        Path(r"C:\Users\18163\Desktop")
        / "\u4e91\u6280_\u8f6f\u4ef6\u8bbe\u8ba1\u8bf4\u660e\u4e66\u6e90\u4ee3\u7801_1.docx"
    )

    parser = argparse.ArgumentParser(
        description="Generate the T003 source-code design DOCX from the selected template."
    )
    parser.add_argument("--repo-root", type=Path, default=repo_root)
    parser.add_argument("--template", type=Path, default=default_template)
    parser.add_argument(
        "--inventory",
        type=Path,
        default=output_dir / "02-\u6e90\u7801\u9009\u53d6\u6e05\u5355.md",
    )
    parser.add_argument("--output", type=Path, default=default_output)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        result = generate(
            template_path=args.template.resolve(),
            inventory_path=args.inventory.resolve(),
            output_path=args.output.resolve(),
            repo_root=args.repo_root.resolve(),
        )
    except Exception as error:
        print(
            json.dumps(
                {"status": "failed", "error": str(error)},
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
