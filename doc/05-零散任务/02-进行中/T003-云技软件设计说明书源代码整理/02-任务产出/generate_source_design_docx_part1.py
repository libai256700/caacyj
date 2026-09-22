from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import uuid
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image


OLD_BRAND = "\u6765\u7ed8"
NEW_BRAND = "\u4e91\u6280"
START_MARKER = "\u7a0b\u5e8f\u5f00\u59cb\uff01\uff01\uff01"
END_MARKER = "\u7a0b\u5e8f\u7ed3\u675f\uff01\uff01\uff01"
MEDIA_PATH = "word/media/image1.png"
EXPECTED_TEMPLATE_SHA256 = (
    "1eb9fef66254c5ec4f894fdf7d005b3446da70b1a94fd5c65ee4a2e00ac2fc82"
)
EXPECTED_PART2_SHA256 = (
    "0530f54d3f1936daaaabb62a59c07fca6201461d9b421904f80911e79a27dcb1"
)
EXPECTED_CODE_PPR_SHA256 = (
    "f4d3882b03d18cf1f78e5736e6657d1de0f356f6239260fa7543147ff83b6aa1"
)
EXPECTED_CODE_RPR_SHA256 = (
    "71036b9ff11a2a284d3801b749f72cef16bc5aa61536bee293b275e94ad6fb7e"
)
EXPECTED_PATHS = (
    "code/develop/yunjikeji/src/main.ts",
    "code/develop/yunjikeji/src/pages/home.vue",
    "code/develop/yunjikeji/src/pages/practice/answer.vue",
    "code/develop/yunjikeji/src/services/practice.ts",
)
HEADER_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"
)
PRESERVED_XML_PARTS = (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
)
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
CREDENTIAL_PATTERNS = (
    ("private-key-header", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    (
        "aws-access-key",
        re.compile(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b", re.IGNORECASE),
    ),
    (
        "long-sk-key",
        re.compile(r"\bsk-[A-Za-z0-9_-]{16,}\b", re.IGNORECASE),
    ),
    (
        "literal-sensitive-assignment",
        re.compile(
            r"(?:password|passwd|pwd|secret|token|api[_-]?key|apikey|"
            r"access[_-]?key|secret[_-]?key|private[_-]?key)\s*[:=]\s*"
            r"[\"'][^\"'\r\n]{4,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "literal-authorization",
        re.compile(
            r"authorization\s*[:=]\s*[\"'](?:bearer|basic)\s+"
            r"[A-Za-z0-9._+/=-]{8,}[\"']",
            re.IGNORECASE,
        ),
    ),
)


@dataclass(frozen=True)
class SourceItem:
    order: int
    project: str
    relative_path: str
    language: str
    expected_lines: int
    expected_chars: int
    expected_sha256: str
    source_path: Path
    source_text: str
    source_lines: tuple[str, ...]

    @property
    def metadata_before(self) -> tuple[str, str, str, str]:
        return (
            f"\u5de5\u7a0b\uff1a{self.project}",
            f"\u6e90\u6587\u4ef6\uff1a{self.relative_path}",
            f"\u8bed\u8a00\uff1a{self.language}",
            "\u6e90\u4ee3\u7801\uff1a",
        )

    @property
    def metadata_after(self) -> str:
        return f"\u6587\u4ef6\u7ed3\u675f\uff1a{self.relative_path}"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def canonical_xml(element_or_bytes) -> bytes:
    element = (
        etree.fromstring(element_or_bytes)
        if isinstance(element_or_bytes, bytes)
        else element_or_bytes
    )
    return etree.tostring(element, method="c14n")


def canonical_sha256(element) -> str:
    return sha256_bytes(canonical_xml(element))


def normalize_newlines(text: str) -> str:
    return re.sub(r"\r\n|\r|\n", "\n", text)


def source_lines(text: str) -> tuple[str, ...]:
    normalized = normalize_newlines(text)
    if normalized == "":
        return ()
    lines = normalized.split("\n")
    if normalized.endswith("\n"):
        lines.pop()
    return tuple(lines)


def utf16_character_count(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def strip_code_ticks(value: str) -> str:
    value = value.strip()
    if value.startswith("`") and value.endswith("`") and len(value) >= 2:
        return value[1:-1]
    return value


def discover_repo_root(script_path: Path) -> Path:
    for candidate in (script_path.parent, *script_path.parents):
        if (candidate / "code" / "develop").is_dir():
            return candidate.resolve()
    raise RuntimeError("Cannot locate repository root")


def is_within(path: Path, root: Path) -> bool:
    try:
        return os.path.commonpath((str(path), str(root))) == str(root)
    except ValueError:
        return False


def validate_xml_characters(text: str, label: str) -> None:
    for index, char in enumerate(text):
        codepoint = ord(char)
        valid = (
            codepoint in (0x09, 0x0A, 0x0D)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        )
        if not valid:
            raise RuntimeError(
                f"Invalid XML character U+{codepoint:04X} in {label} at {index}"
            )


def load_inventory(inventory_path: Path, repo_root: Path) -> list[SourceItem]:
    inventory_text = inventory_path.read_text(encoding="utf-8")
    items: list[SourceItem] = []
    for line in inventory_text.splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 9 or not cells[0].isdigit():
            continue

        order = int(cells[0])
        project = strip_code_ticks(cells[1])
        relative_path = strip_code_ticks(cells[2]).replace("\\", "/")
        language = strip_code_ticks(cells[3])
        expected_lines = int(cells[5].replace(",", ""))
        expected_chars = int(cells[6].replace(",", ""))
        expected_sha256 = strip_code_ticks(cells[7]).lower()

        source_path = (repo_root / relative_path).resolve()
        if not is_within(source_path, repo_root) or not source_path.is_file():
            raise RuntimeError(f"Invalid source path: {relative_path}")
        raw = source_path.read_bytes()
        text = raw.decode("utf-8-sig", errors="strict")
        lines = source_lines(text)
        validate_xml_characters(text, relative_path)

        if sha256_bytes(raw) != expected_sha256:
            raise RuntimeError(f"SHA-256 drift: {relative_path}")
        if len(lines) != expected_lines:
            raise RuntimeError(f"Line-count drift: {relative_path}")
        if utf16_character_count(text) != expected_chars:
            raise RuntimeError(f"Character-count drift: {relative_path}")

        items.append(
            SourceItem(
                order=order,
                project=project,
                relative_path=relative_path,
                language=language,
                expected_lines=expected_lines,
                expected_chars=expected_chars,
                expected_sha256=expected_sha256,
                source_path=source_path,
                source_text=text,
                source_lines=lines,
            )
        )

    items.sort(key=lambda item: item.order)
    if len(items) != 4:
        raise RuntimeError(f"Expected 4 source files, found {len(items)}")
    if tuple(item.relative_path for item in items) != EXPECTED_PATHS:
        raise RuntimeError("Source inventory order or paths do not match the task")
    if any(item.project != "yunjikeji" for item in items):
        raise RuntimeError("All part-1 sources must belong to yunjikeji")
    if sum(item.expected_lines for item in items) != 2_624:
        raise RuntimeError("Source line total is not 2,624")
    if sum(item.expected_chars for item in items) != 69_824:
        raise RuntimeError("Source character total is not 69,824")
    return items


def scan_credentials(items: list[SourceItem]) -> list[dict]:
    hits: list[dict] = []
    for item in items:
        for line_number, line in enumerate(item.source_lines, start=1):
            for pattern_name, pattern in CREDENTIAL_PATTERNS:
                if pattern.search(line):
                    hits.append(
                        {
                            "path": item.relative_path,
                            "line": line_number,
                            "pattern": pattern_name,
                        }
                    )
    return hits


def git_evidence(repo_root: Path) -> dict[str, str]:
    commands = {
        "status": ["git", "status", "--short", "--", "code/develop"],
        "diff": ["git", "diff", "--name-only", "--", "code/develop"],
        "cached_diff": [
            "git",
            "diff",
            "--cached",
            "--name-only",
            "--",
            "code/develop",
        ],
    }
    evidence: dict[str, str] = {}
    for name, command in commands.items():
        result = subprocess.run(
            command,
            cwd=repo_root,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        evidence[name] = result.stdout.decode("utf-8", errors="strict").strip()
    if any(evidence.values()):
        raise RuntimeError(f"Business source Git check failed: {evidence}")
    return evidence


def replace_equal_length_text_nodes(root, old: str, new: str) -> int:
    if len(old) != len(new):
        raise ValueError("Brand replacement must preserve text length")
    nodes = root.xpath(".//w:t")
    original = [node.text or "" for node in nodes]
    combined = "".join(original)
    count = combined.count(old)
    replaced = combined.replace(old, new)
    offset = 0
    for node, original_text in zip(nodes, original):
        node.text = replaced[offset : offset + len(original_text)]
        offset += len(original_text)
    return count


def replace_header_brand(document: Document) -> int:
    count = 0
    seen: set[str] = set()
    for part in document.part.package.parts:
        if part.content_type != HEADER_CONTENT_TYPE:
            continue
        part_name = str(part.partname)
        if part_name in seen:
            continue
        seen.add(part_name)
        count += replace_equal_length_text_nodes(part.element, OLD_BRAND, NEW_BRAND)
    if count < 1:
        raise RuntimeError("Old brand was not found in editable header text")
    return count


def append_text_with_tabs(run_element, text: str) -> None:
    for token in re.split(r"(\t)", text):
        if token == "\t":
            run_element.append(OxmlElement("w:tab"))
        elif token:
            text_element = OxmlElement("w:t")
            if token[0].isspace() or token[-1].isspace():
                text_element.set(XML_SPACE, "preserve")
            text_element.text = token
            run_element.append(text_element)


def build_line_paragraph(text: str, ppr, rpr):
    paragraph = OxmlElement("w:p")
    paragraph.append(deepcopy(ppr))
    run = OxmlElement("w:r")
    run.append(deepcopy(rpr))
    append_text_with_tabs(run, text)
    paragraph.append(run)
    return paragraph


def rebuild_body(document: Document, items: list[SourceItem]) -> dict:
    paragraphs = document.paragraphs
    if len(paragraphs) != 3_816:
        raise RuntimeError("Template body paragraph count is not 3,816")
    if paragraphs[0].text != START_MARKER:
        raise RuntimeError("Template opening marker is incorrect")
    end_indexes = [i for i, paragraph in enumerate(paragraphs) if paragraph.text == END_MARKER]
    if end_indexes != [3_814]:
        raise RuntimeError(f"Unexpected template closing marker indexes: {end_indexes}")

    prototype = paragraphs[1]
    if prototype.text != "<template>" or not prototype.runs:
        raise RuntimeError("Representative template code paragraph is unavailable")
    code_ppr = prototype._p.pPr
    code_rpr = prototype.runs[0]._r.rPr
    if canonical_sha256(code_ppr) != EXPECTED_CODE_PPR_SHA256:
        raise RuntimeError("Representative code pPr changed")
    if canonical_sha256(code_rpr) != EXPECTED_CODE_RPR_SHA256:
        raise RuntimeError("Representative code rPr changed")

    opening = deepcopy(paragraphs[0]._p)
    closing = deepcopy(paragraphs[3_814]._p)
    body = document._body._element
    section_properties = body.sectPr
    if section_properties is None:
        raise RuntimeError("Template section properties are missing")
    for child in list(body):
        if child is not section_properties:
            body.remove(child)

    body.insert(len(body) - 1, opening)
    source_paragraphs = 0
    metadata_paragraphs = 0
    for item in items:
        for metadata in item.metadata_before:
            body.insert(
                len(body) - 1,
                build_line_paragraph(metadata, code_ppr, code_rpr),
            )
            metadata_paragraphs += 1
        for line in item.source_lines:
            body.insert(len(body) - 1, build_line_paragraph(line, code_ppr, code_rpr))
            source_paragraphs += 1
        body.insert(
            len(body) - 1,
            build_line_paragraph(item.metadata_after, code_ppr, code_rpr),
        )
        metadata_paragraphs += 1
    body.insert(len(body) - 1, closing)

    expected_total = 2 + metadata_paragraphs + source_paragraphs
    if expected_total != 2_646 or len(document.paragraphs) != expected_total:
        raise RuntimeError("Generated body paragraph count is not 2,646")
    return {
        "template_paragraphs": len(paragraphs),
        "output_paragraphs": expected_total,
        "source_line_paragraphs": source_paragraphs,
        "metadata_paragraphs": metadata_paragraphs,
        "code_ppr_sha256": EXPECTED_CODE_PPR_SHA256,
        "code_rpr_sha256": EXPECTED_CODE_RPR_SHA256,
    }


def create_white_png() -> bytes:
    image = Image.new("RGB", (640, 960), (255, 255, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def rebuild_media(source_docx: Path, target_docx: Path, media: bytes) -> None:
    found = False
    with zipfile.ZipFile(source_docx, "r") as source_zip:
        with zipfile.ZipFile(target_docx, "w") as target_zip:
            target_zip.comment = source_zip.comment
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == MEDIA_PATH:
                    data = media
                    found = True
                target_zip.writestr(info, data)
    if not found:
        raise RuntimeError(f"Template media not found: {MEDIA_PATH}")


def replace_brand_in_lxml_header(root, old: str, new: str) -> int:
    nodes = root.xpath(".//w:t", namespaces=NS)
    original = [node.text or "" for node in nodes]
    combined = "".join(original)
    count = combined.count(old)
    replaced = combined.replace(old, new)
    offset = 0
    for node, original_text in zip(nodes, original):
        node.text = replaced[offset : offset + len(original_text)]
        offset += len(original_text)
    return count


def count_page_fields(header_root) -> int:
    instructions = header_root.xpath(".//w:instrText/text()", namespaces=NS)
    simple_fields = header_root.xpath(".//w:fldSimple/@w:instr", namespaces=NS)
    return sum(bool(re.search(r"\bPAGE\b", value, re.IGNORECASE)) for value in (*instructions, *simple_fields))


def package_validation(template_path: Path, output_path: Path) -> dict:
    with zipfile.ZipFile(template_path, "r") as template_zip:
        template_parts = {name for name in template_zip.namelist() if not name.endswith("/")}
        template_document = etree.fromstring(template_zip.read("word/document.xml"))
        template_sections = template_document.xpath(
            "/w:document/w:body/w:sectPr", namespaces=NS
        )
        if len(template_sections) != 1:
            raise RuntimeError("Template must contain one section")
        template_section_xml = canonical_xml(template_sections[0])
        preserved_parts = {
            name: canonical_xml(template_zip.read(name)) for name in PRESERVED_XML_PARTS
        }
        template_headers = sorted(
            name
            for name in template_parts
            if name.startswith("word/header") and name.endswith(".xml")
        )
        template_header_rels = sorted(
            name
            for name in template_parts
            if name.startswith("word/_rels/header") and name.endswith(".rels")
        )
        expected_headers: dict[str, bytes] = {}
        template_page_fields = 0
        for name in template_headers:
            root = etree.fromstring(template_zip.read(name))
            template_page_fields += count_page_fields(root)
            replace_brand_in_lxml_header(root, OLD_BRAND, NEW_BRAND)
            expected_headers[name] = canonical_xml(root)
        expected_header_rels = {
            name: canonical_xml(template_zip.read(name)) for name in template_header_rels
        }

    with zipfile.ZipFile(output_path, "r") as output_zip:
        if output_zip.testzip() is not None:
            raise RuntimeError("Output ZIP CRC check failed")
        output_parts = {name for name in output_zip.namelist() if not name.endswith("/")}
        if output_parts != template_parts:
            raise RuntimeError("Output actual package parts differ from template")

        xml_names = sorted(
            name
            for name in output_parts
            if name.endswith(".xml") or name.endswith(".rels")
        )
        parsed = {name: etree.fromstring(output_zip.read(name)) for name in xml_names}
        document_root = parsed["word/document.xml"]
        body_paragraphs = document_root.xpath(
            "/w:document/w:body/w:p", namespaces=NS
        )
        tables = document_root.xpath(".//w:tbl", namespaces=NS)
        sections = document_root.xpath(".//w:sectPr", namespaces=NS)
        if len(body_paragraphs) != 2_646 or tables or len(sections) != 1:
            raise RuntimeError("Output XML is not 2,646 paragraphs / 0 tables / 1 section")
        if canonical_xml(sections[0]) != template_section_xml:
            raise RuntimeError("Section properties changed")

        for name, expected in preserved_parts.items():
            if canonical_xml(output_zip.read(name)) != expected:
                raise RuntimeError(f"Preserved style part changed: {name}")
        for name, expected in expected_headers.items():
            if canonical_xml(output_zip.read(name)) != expected:
                raise RuntimeError(f"Header structure changed beyond brand text: {name}")
        for name, expected in expected_header_rels.items():
            if canonical_xml(output_zip.read(name)) != expected:
                raise RuntimeError(f"Header relationship changed: {name}")

        output_page_fields = sum(count_page_fields(parsed[name]) for name in template_headers)
        if template_page_fields != 2 or output_page_fields != template_page_fields:
            raise RuntimeError("The two PAGE fields were not preserved")

        all_xml_text = "\n".join("".join(root.itertext()) for root in parsed.values())
        header_text = "\n".join(
            "".join(parsed[name].itertext()) for name in template_headers
        )
        if OLD_BRAND in all_xml_text or NEW_BRAND not in header_text:
            raise RuntimeError("Header brand replacement is incomplete")

        media_names = sorted(
            name for name in output_parts if name.startswith("word/media/")
        )
        if media_names != [MEDIA_PATH]:
            raise RuntimeError(f"Expected one media file, found {media_names}")
        media_bytes = output_zip.read(MEDIA_PATH)
        with Image.open(io.BytesIO(media_bytes)) as media_image:
            media_image.load()
            if media_image.size != (640, 960):
                raise RuntimeError("Replacement media size is incorrect")
            extrema = media_image.convert("RGB").getextrema()
            if extrema != ((255, 255), (255, 255), (255, 255)):
                raise RuntimeError("Replacement media is not pure white")

        tab_count = len(document_root.xpath(".//w:tab", namespaces=NS))
        line_break_count = len(document_root.xpath(".//w:br", namespaces=NS))

    return {
        "zip_crc": "ok",
        "package_file_parts": len(output_parts),
        "xml_parts_parsed": len(xml_names),
        "paragraphs": len(body_paragraphs),
        "tables": len(tables),
        "sections": len(sections),
        "media_count": len(media_names),
        "media_path": MEDIA_PATH,
        "media_size": [640, 960],
        "media_pure_white": True,
        "header_xml": template_headers,
        "header_relationships_preserved": True,
        "page_fields": output_page_fields,
        "old_brand_occurrences": 0,
        "new_brand_in_header": True,
        "style_parts_preserved": list(PRESERVED_XML_PARTS),
        "section_properties_preserved": True,
        "tab_elements": tab_count,
        "body_line_break_elements": line_break_count,
    }


def section_values(document: Document) -> dict[str, int]:
    section = document.sections[0]
    return {
        "page_width": int(section.page_width),
        "page_height": int(section.page_height),
        "top_margin": int(section.top_margin),
        "bottom_margin": int(section.bottom_margin),
        "left_margin": int(section.left_margin),
        "right_margin": int(section.right_margin),
        "header_distance": int(section.header_distance),
        "footer_distance": int(section.footer_distance),
    }


def to_millimeters(values: dict[str, int]) -> dict[str, float]:
    return {name: round(value / 36_000, 3) for name, value in values.items()}


def document_validation(
    template_path: Path,
    output_path: Path,
    items: list[SourceItem],
) -> dict:
    template = Document(template_path)
    output = Document(output_path)
    if len(output.paragraphs) != 2_646 or len(output.tables) != 0 or len(output.sections) != 1:
        raise RuntimeError("Reopened output structure is incorrect")
    if output.paragraphs[0].text != START_MARKER or output.paragraphs[-1].text != END_MARKER:
        raise RuntimeError("Reopened output boundary markers are incorrect")
    template_section_values = section_values(template)
    output_section_values = section_values(output)
    if output_section_values != template_section_values:
        raise RuntimeError("A4 page size or margins changed")

    prototype_ppr = template.paragraphs[1]._p.pPr
    prototype_rpr = template.paragraphs[1].runs[0]._r.rPr
    prototype_ppr_xml = canonical_xml(prototype_ppr)
    prototype_rpr_xml = canonical_xml(prototype_rpr)
    for paragraph in output.paragraphs[1:-1]:
        if canonical_xml(paragraph._p.pPr) != prototype_ppr_xml:
            raise RuntimeError("Generated line paragraph pPr differs from A prototype")
        if len(paragraph.runs) != 1:
            raise RuntimeError("Generated line paragraph does not contain exactly one run")
        if canonical_xml(paragraph.runs[0]._r.rPr) != prototype_rpr_xml:
            raise RuntimeError("Generated line paragraph rPr differs from A prototype")

    cursor = 1
    reconstructed: dict[str, list[str]] = {}
    metadata_count = 0
    for item in items:
        for expected_metadata in item.metadata_before:
            if output.paragraphs[cursor].text != expected_metadata:
                raise RuntimeError(
                    f"Metadata mismatch for {item.relative_path} at paragraph {cursor}"
                )
            cursor += 1
            metadata_count += 1
        extracted_lines = [
            paragraph.text
            for paragraph in output.paragraphs[
                cursor : cursor + item.expected_lines
            ]
        ]
        if tuple(extracted_lines) != item.source_lines:
            raise RuntimeError(f"Source line sequence mismatch: {item.relative_path}")
        reconstructed[item.relative_path] = extracted_lines
        cursor += item.expected_lines
        if output.paragraphs[cursor].text != item.metadata_after:
            raise RuntimeError(f"File-end metadata mismatch: {item.relative_path}")
        cursor += 1
        metadata_count += 1
    if cursor != len(output.paragraphs) - 1:
        raise RuntimeError("Unexpected paragraphs remain before the closing marker")

    expected_tabs = sum(line.count("\t") for item in items for line in item.source_lines)
    actual_tabs = sum(
        paragraph.text.count("\t") for paragraph in output.paragraphs[1:-1]
    )
    if actual_tabs != expected_tabs:
        raise RuntimeError("Tab indentation count changed")

    generated_text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    credential_hits = [
        name for name, pattern in CREDENTIAL_PATTERNS if pattern.search(generated_text)
    ]
    if credential_hits:
        raise RuntimeError(f"Credential patterns found in output: {credential_hits}")

    return {
        "reopen": "ok",
        "paragraphs": len(output.paragraphs),
        "tables": len(output.tables),
        "sections": len(output.sections),
        "files_reconstructed": len(reconstructed),
        "source_lines_reconstructed": sum(len(lines) for lines in reconstructed.values()),
        "metadata_paragraphs": metadata_count,
        "line_sequences_exact": True,
        "empty_source_lines_preserved": sum(
            line == "" for item in items for line in item.source_lines
        ),
        "tab_characters_preserved": actual_tabs,
        "code_ppr_preserved": True,
        "code_rpr_preserved": True,
        "page_values_emu": output_section_values,
        "page_values_mm": to_millimeters(output_section_values),
        "credential_hits": credential_hits,
    }


def generate(
    template_path: Path,
    inventory_path: Path,
    output_path: Path,
    part2_path: Path,
    repo_root: Path,
) -> dict:
    if not template_path.is_file() or not inventory_path.is_file():
        raise FileNotFoundError("Template or inventory is missing")
    if not part2_path.is_file():
        raise FileNotFoundError(part2_path)
    if sha256_file(template_path) != EXPECTED_TEMPLATE_SHA256:
        raise RuntimeError("Template A SHA-256 does not match the approved input")
    part2_sha_before = sha256_file(part2_path)
    if part2_sha_before != EXPECTED_PART2_SHA256:
        raise RuntimeError("Existing _1 output SHA-256 changed before generation")

    items = load_inventory(inventory_path, repo_root)
    credential_hits = scan_credentials(items)
    if credential_hits:
        raise RuntimeError(f"Credential-shaped values found: {credential_hits}")
    git_before = git_evidence(repo_root)

    document = Document(template_path)
    if len(document.tables) != 0 or len(document.sections) != 1:
        raise RuntimeError("Template A must contain 0 tables and 1 section")
    brand_replacements = replace_header_brand(document)
    body_result = rebuild_body(document, items)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    unique_id = uuid.uuid4().hex
    python_docx_path = output_path.parent / (
        f".{output_path.stem}.{unique_id}.python-docx.docx"
    )
    candidate_path = output_path.parent / (
        f".{output_path.stem}.{unique_id}.candidate.docx"
    )
    try:
        document.save(python_docx_path)
        rebuild_media(python_docx_path, candidate_path, create_white_png())
        package_result = package_validation(template_path, candidate_path)
        document_result = document_validation(template_path, candidate_path, items)

        expected_tabs = sum(line.count("\t") for item in items for line in item.source_lines)
        if package_result["tab_elements"] != expected_tabs:
            raise RuntimeError("w:tab element count does not match source indentation")
        if package_result["body_line_break_elements"] != 0:
            raise RuntimeError("Part-1 body must not use in-paragraph line breaks")

        if sha256_file(template_path) != EXPECTED_TEMPLATE_SHA256:
            raise RuntimeError("Template A changed during generation")
        if sha256_file(part2_path) != part2_sha_before:
            raise RuntimeError("Existing _1 output changed during generation")
        for item in items:
            if sha256_file(item.source_path) != item.expected_sha256:
                raise RuntimeError(f"Source changed during generation: {item.relative_path}")
        git_after = git_evidence(repo_root)

        candidate_sha256 = sha256_file(candidate_path)
        candidate_size = candidate_path.stat().st_size
        os.replace(candidate_path, output_path)
        if sha256_file(output_path) != candidate_sha256:
            raise RuntimeError("Final output differs from validated candidate")
        Document(output_path)
        if sha256_file(part2_path) != EXPECTED_PART2_SHA256:
            raise RuntimeError("Existing _1 output changed after final placement")

        return {
            "status": "passed",
            "template": str(template_path),
            "template_size": template_path.stat().st_size,
            "template_sha256": EXPECTED_TEMPLATE_SHA256,
            "inventory": str(inventory_path),
            "output": str(output_path),
            "output_size": candidate_size,
            "output_sha256": candidate_sha256,
            "part2_output": str(part2_path),
            "part2_sha256_before": part2_sha_before,
            "part2_sha256_after": sha256_file(part2_path),
            "source_files": len(items),
            "source_lines": sum(item.expected_lines for item in items),
            "source_characters": sum(item.expected_chars for item in items),
            "header_brand_replacements": brand_replacements,
            "credential_hits": credential_hits,
            "git_before": git_before,
            "git_after": git_after,
            "body": body_result,
            "package": package_result,
            "document": document_result,
        }
    finally:
        python_docx_path.unlink(missing_ok=True)
        candidate_path.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    script_path = Path(__file__).resolve()
    repo_root = discover_repo_root(script_path)
    output_dir = script_path.parent
    template = (
        Path(r"C:\Users\18163\Downloads")
        / "\u98de\u4e6620260724-105710"
        / "\u6765\u7ed8_\u8f6f\u4ef6\u8bbe\u8ba1\u8bf4\u660e\u4e66\u6e90\u4ee3\u7801.docx"
    )
    desktop = Path(r"C:\Users\18163\Desktop")
    parser = argparse.ArgumentParser(
        description="Generate the line-per-paragraph T003 source-code DOCX."
    )
    parser.add_argument("--repo-root", type=Path, default=repo_root)
    parser.add_argument("--template", type=Path, default=template)
    parser.add_argument(
        "--inventory",
        type=Path,
        default=output_dir / "05-\u65e0\u540e\u7f00\u7248\u6e90\u7801\u6e05\u5355.md",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=desktop
        / "\u4e91\u6280_\u8f6f\u4ef6\u8bbe\u8ba1\u8bf4\u660e\u4e66\u6e90\u4ee3\u7801.docx",
    )
    parser.add_argument(
        "--part2",
        type=Path,
        default=desktop
        / "\u4e91\u6280_\u8f6f\u4ef6\u8bbe\u8ba1\u8bf4\u660e\u4e66\u6e90\u4ee3\u7801_1.docx",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        result = generate(
            template_path=args.template.resolve(),
            inventory_path=args.inventory.resolve(),
            output_path=args.output.resolve(),
            part2_path=args.part2.resolve(),
            repo_root=args.repo_root.resolve(),
        )
    except Exception as error:
        print(
            json.dumps(
                {"status": "failed", "error": str(error)},
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
