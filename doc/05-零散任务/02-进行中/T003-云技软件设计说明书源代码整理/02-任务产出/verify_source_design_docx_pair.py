from __future__ import annotations

import ast
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W_NS, "wp": WP_NS}

EXPECTED_PROJECTS = {
    "yunjikeji",
    "yunjikeji-admin-ui",
    "yunjikeji-server",
    "yunjikeji-admin-server",
    "fly-llm",
}

EXPECTED = {
    "A": {
        "output_name": "云技_软件设计说明书源代码.docx",
        "output_sha256": "731226e7db3ac32a8f5a0b3f3a29ce6ea78d1fc2487607c1ea693ed0e4ad88ef",
        "output_bytes": 54286,
        "reference_name": "来绘_软件设计说明书源代码.docx",
        "reference_sha256": "1eb9fef66254c5ec4f894fdf7d005b3446da70b1a94fd5c65ee4a2e00ac2fc82",
        "reference_bytes": 211878,
        "inventory_name": "05-无后缀版源码清单.md",
        "inventory_sha256": "27b0510d4b0c80e1ab245911fb9e312beac7c455940bb450274c71be88bf026d",
        "generator_name": "generate_source_design_docx_part1.py",
        "generator_sha256": "efb156af403bc960d0c38b3b34661908da845d95035ac7fd371b5f7ef928f848",
        "record_name": "06-补充生成记录.md",
        "source_files": 4,
        "source_lines": 2624,
        "source_chars": 69824,
        "paragraphs": 2646,
        "package_files": 17,
        "xml_parts": 16,
        "page_fields": 2,
        "drawings": 2,
        "picts": 1,
        "breaks": 0,
    },
    "B": {
        "output_name": "云技_软件设计说明书源代码_1.docx",
        "output_sha256": "0530f54d3f1936daaaabb62a59c07fca6201461d9b421904f80911e79a27dcb1",
        "output_bytes": 80660,
        "reference_name": "来绘_软件设计说明书源代码_1.docx",
        "reference_sha256": "3271e0e1a64ad834808723230995948df5b23f9988adf8c0dd96a4b2f16301b2",
        "reference_bytes": 144634,
        "inventory_name": "02-源码选取清单.md",
        "inventory_sha256": "c8eb4956fecc869ca68d82a4849909bdcbf81101af57e3dd5224afc6b8bd940e",
        "generator_name": "generate_source_design_docx.py",
        "generator_sha256": "dc0676a0b75d8c01f2be76f2102e1ea578445834fc94573ad2db422897cb71cf",
        "record_name": "03-生成记录.md",
        "source_files": 18,
        "source_lines": 8329,
        "source_chars": 310972,
        "paragraphs": 20,
        "package_files": 15,
        "xml_parts": 14,
        "page_fields": 0,
        "drawings": 1,
        "picts": 0,
        "breaks": 8401,
    },
}

PRESERVED_PARTS = (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
    "word/_rels/document.xml.rels",
    "word/_rels/header1.xml.rels",
)
REQUIRED_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/_rels/document.xml.rels",
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
    "word/settings.xml",
    "word/header1.xml",
    "word/_rels/header1.xml.rels",
    "word/media/image1.png",
}
EXCLUDED_PATH_MARKERS = (
    "/.env",
    "/tmp/",
    "/node_modules/",
    "/dist/",
    "/target/",
    "/unpackage/",
    "application.yml",
    "application.yaml",
    "application-",
    "账号",
    "密码",
)
EXCLUDED_SUFFIXES = {
    ".pem",
    ".key",
    ".p12",
    ".pfx",
    ".jks",
    ".crt",
    ".cer",
    ".sql",
}
CREDENTIAL_PATTERNS = (
    ("private-key", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    ("certificate", re.compile(r"-----BEGIN CERTIFICATE-----")),
    ("aws-access-key", re.compile(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b", re.IGNORECASE)),
    ("long-sk-key", re.compile(r"\bsk-[A-Za-z0-9_-]{16,}\b", re.IGNORECASE)),
    (
        "literal-sensitive-assignment",
        re.compile(
            r"\b(?:password|passwd|pwd|secret|token|api[_-]?key|apikey|"
            r"access[_-]?key|secret[_-]?key|private[_-]?key)\b\s*[:=]\s*"
            r"[\"'][^\"'\r\n]{4,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "literal-authorization",
        re.compile(
            r"\bauthorization\b\s*[:=]\s*[\"'](?:bearer|basic)\s+"
            r"[A-Za-z0-9._+/=-]{8,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "credential-in-url",
        re.compile(r"\b[a-z][a-z0-9+.-]*://[^\s/:]+:[^\s/@]+@", re.IGNORECASE),
    ),
)


@dataclass(frozen=True)
class SourceItem:
    order: int
    project: str
    relative_path: str
    language: str
    expected_lines: int
    expected_chars: int
    expected_sha256: str
    source_text: str

    @property
    def normalized_text(self) -> str:
        return normalize_newlines(self.source_text)

    @property
    def container_text(self) -> str:
        return (
            f"工程：{self.project}\n"
            f"路径：{self.relative_path}\n"
            f"语言：{self.language}\n"
            "源码：\n"
            + self.normalized_text
        )


@dataclass(frozen=True)
class Check:
    group: str
    name: str
    passed: bool
    detail: str


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def normalize_newlines(text: str) -> str:
    return re.sub(r"\r\n|\r|\n", "\n", text)


def utf16_character_count(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def strip_ticks(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value.startswith("`") and value.endswith("`"):
        return value[1:-1]
    return value


def find_repo_root(start: Path) -> Path:
    for candidate in (start.parent, *start.parents):
        if (candidate / ".git").exists() and (candidate / "code" / "develop").is_dir():
            return candidate.resolve()
    raise RuntimeError("Cannot locate repository root")


def is_within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def load_inventory(path: Path, repo_root: Path) -> list[SourceItem]:
    items: list[SourceItem] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 9 or not cells[0].isdigit():
            continue
        relative_path = strip_ticks(cells[2]).replace("\\", "/")
        source_path = (repo_root / relative_path).resolve()
        if not relative_path.startswith("code/develop/") or not is_within(source_path, repo_root):
            raise RuntimeError(f"Inventory path outside repository: {relative_path}")
        raw = source_path.read_bytes()
        items.append(
            SourceItem(
                order=int(cells[0]),
                project=strip_ticks(cells[1]),
                relative_path=relative_path,
                language=strip_ticks(cells[3]),
                expected_lines=int(cells[5].replace(",", "")),
                expected_chars=int(cells[6].replace(",", "")),
                expected_sha256=strip_ticks(cells[7]).lower(),
                source_text=raw.decode("utf-8-sig", errors="strict"),
            )
        )
    return items


def load_package(path: Path) -> tuple[dict[str, bytes], str | None]:
    with zipfile.ZipFile(path) as archive:
        crc_error = archive.testzip()
        package = {name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}
    return package, crc_error


def parse_xml_parts(package: dict[str, bytes]) -> dict[str, etree._Element]:
    return {
        name: etree.fromstring(data)
        for name, data in package.items()
        if name.endswith((".xml", ".rels"))
    }


def c14n(data_or_element: bytes | etree._Element | None) -> bytes:
    if data_or_element is None:
        return b"<none>"
    element = (
        etree.fromstring(data_or_element)
        if isinstance(data_or_element, bytes)
        else data_or_element
    )
    return etree.tostring(element, method="c14n", with_comments=False)


def paragraph_properties(root: etree._Element) -> list[bytes]:
    paragraphs = root.xpath("/w:document/w:body/w:p", namespaces=NS)
    return [c14n(paragraph.find(f"{{{W_NS}}}pPr")) for paragraph in paragraphs]


def first_run_properties(paragraph: etree._Element) -> bytes:
    run = paragraph.find(f"{{{W_NS}}}r")
    return c14n(run.find(f"{{{W_NS}}}rPr") if run is not None else None)


def page_metrics(document: Document) -> dict[str, float]:
    section = document.sections[0]
    to_mm = lambda value: round(float(value) / 36000.0, 3)
    return {
        "width_mm": to_mm(section.page_width),
        "height_mm": to_mm(section.page_height),
        "top_mm": to_mm(section.top_margin),
        "bottom_mm": to_mm(section.bottom_margin),
        "left_mm": to_mm(section.left_margin),
        "right_mm": to_mm(section.right_margin),
        "header_mm": to_mm(section.header_distance),
        "footer_mm": to_mm(section.footer_distance),
    }


def page_metrics_valid(metrics: dict[str, float]) -> bool:
    expected = {
        "width_mm": 210.0,
        "height_mm": 297.0,
        "top_mm": 10.0,
        "bottom_mm": 10.0,
        "left_mm": 31.75,
        "right_mm": 31.75,
        "header_mm": 12.7,
        "footer_mm": 12.7,
    }
    return all(abs(metrics[key] - value) <= 0.2 for key, value in expected.items())


def scan_credentials(text: str) -> list[str]:
    return [name for name, pattern in CREDENTIAL_PATTERNS if pattern.search(text)]


def imported_modules(tree: ast.AST) -> set[str]:
    modules: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            modules.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            modules.add(node.module.split(".")[0])
    return modules


def called_names(tree: ast.AST) -> set[str]:
    names: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            if isinstance(node.func, ast.Name):
                names.add(node.func.id)
            elif isinstance(node.func, ast.Attribute):
                names.add(node.func.attr)
    return names


def run_git(repo_root: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["git", *args],
        cwd=repo_root,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )


def detect_office_renderers() -> list[str]:
    found: list[str] = []
    for executable in ("WINWORD.EXE", "wps.exe", "soffice.exe", "libreoffice.exe"):
        resolved = shutil.which(executable)
        if resolved:
            found.append(resolved)
    for path in (
        Path(r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE"),
        Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE"),
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ):
        if path.is_file():
            found.append(str(path))
    return sorted(set(found))


def main() -> int:
    script_path = Path(__file__).resolve()
    output_dir = script_path.parent
    repo_root = find_repo_root(script_path)
    reference_dir = Path.home() / "Downloads" / "飞书20260724-105710"
    desktop = Path.home() / "Desktop"
    checks: list[Check] = []

    def record(group: str, name: str, passed: bool, detail: str) -> None:
        checks.append(Check(group, name, bool(passed), detail))

    expected_output_names = sorted(spec["output_name"] for spec in EXPECTED.values())
    delivered_names = sorted(path.name for path in desktop.glob("云技_软件设计说明书源代码*.docx"))
    record(
        "E2E",
        "two exact desktop deliverables",
        delivered_names == expected_output_names,
        repr(delivered_names),
    )

    state: dict[str, dict[str, object]] = {}
    combined_projects: set[str] = set()

    for variant, spec in EXPECTED.items():
        output_path = desktop / str(spec["output_name"])
        reference_path = reference_dir / str(spec["reference_name"])
        inventory_path = output_dir / str(spec["inventory_name"])
        generator_path = output_dir / str(spec["generator_name"])
        generation_record_path = output_dir / str(spec["record_name"])

        actual_output_hash = sha256_file(output_path)
        actual_output_bytes = output_path.stat().st_size
        actual_reference_hash = sha256_file(reference_path)
        actual_reference_bytes = reference_path.stat().st_size
        actual_inventory_hash = sha256_file(inventory_path)
        actual_generator_hash = sha256_file(generator_path)
        generation_record_text = generation_record_path.read_text(encoding="utf-8")

        record(
            "E2E",
            f"{variant} artifact hash and size",
            actual_output_hash == spec["output_sha256"]
            and actual_output_bytes == spec["output_bytes"],
            f"sha256={actual_output_hash}, bytes={actual_output_bytes}",
        )
        record(
            "API",
            f"{variant} reference identity",
            actual_reference_hash == spec["reference_sha256"]
            and actual_reference_bytes == spec["reference_bytes"],
            f"sha256={actual_reference_hash}, bytes={actual_reference_bytes}",
        )

        items = load_inventory(inventory_path, repo_root)
        projects = {item.project for item in items}
        combined_projects.update(projects)
        orders = [item.order for item in items]
        missing: list[str] = []
        source_hash_mismatches: list[str] = []
        source_stat_mismatches: list[str] = []
        excluded_paths: list[str] = []
        source_lines = 0
        source_chars = 0
        source_empty_lines = 0
        source_tabs = 0

        for item in items:
            source_path = repo_root / item.relative_path
            if not source_path.is_file():
                missing.append(item.relative_path)
                continue
            raw = source_path.read_bytes()
            if sha256_bytes(raw) != item.expected_sha256:
                source_hash_mismatches.append(item.relative_path)
            lines = item.source_text.splitlines()
            actual_lines = len(lines)
            actual_chars = utf16_character_count(item.source_text)
            source_lines += actual_lines
            source_chars += actual_chars
            source_empty_lines += sum(line == "" for line in lines)
            source_tabs += item.source_text.count("\t")
            if (actual_lines, actual_chars) != (item.expected_lines, item.expected_chars):
                source_stat_mismatches.append(item.relative_path)
            lowered = "/" + item.relative_path.lower()
            if any(marker in lowered for marker in EXCLUDED_PATH_MARKERS) or source_path.suffix.lower() in EXCLUDED_SUFFIXES:
                excluded_paths.append(item.relative_path)

        record(
            "DATA",
            f"{variant} inventory order and count",
            len(items) == spec["source_files"]
            and orders == list(range(1, int(spec["source_files"]) + 1)),
            f"files={len(items)}, orders={orders}",
        )
        record("DATA", f"{variant} all source paths exist", not missing, repr(missing))
        record(
            "DATA",
            f"{variant} source hashes and metadata",
            not source_hash_mismatches
            and not source_stat_mismatches
            and source_lines == spec["source_lines"]
            and source_chars == spec["source_chars"],
            f"lines={source_lines}, chars={source_chars}, hash_mismatch={source_hash_mismatches}, stat_mismatch={source_stat_mismatches}",
        )
        record("DATA", f"{variant} excluded-path policy", not excluded_paths, repr(excluded_paths))

        package, crc_error = load_package(output_path)
        reference_package, reference_crc_error = load_package(reference_path)
        parsed_xml = parse_xml_parts(package)
        reference_xml = parse_xml_parts(reference_package)
        document = Document(output_path)
        reference_document = Document(reference_path)
        document_root = parsed_xml["word/document.xml"]
        reference_document_root = reference_xml["word/document.xml"]

        record(
            "API",
            f"{variant} ZIP and XML integrity",
            crc_error is None
            and reference_crc_error is None
            and len(package) == spec["package_files"]
            and len(parsed_xml) == spec["xml_parts"]
            and REQUIRED_PARTS.issubset(package),
            f"crc={crc_error}, files={len(package)}, xml={len(parsed_xml)}, missing={sorted(REQUIRED_PARTS - package.keys())}",
        )
        record(
            "API",
            f"{variant} package maps to matching reference",
            set(package) == set(reference_package),
            f"output_only={sorted(set(package) - set(reference_package))}, reference_only={sorted(set(reference_package) - set(package))}",
        )
        record(
            "API",
            f"{variant} python-docx structure",
            len(document.paragraphs) == spec["paragraphs"]
            and len(document.tables) == 0
            and len(document.sections) == 1
            and len(reference_document.tables) == 0
            and len(reference_document.sections) == 1,
            f"paragraphs={len(document.paragraphs)}, tables={len(document.tables)}, sections={len(document.sections)}",
        )

        preserved_mismatches = [
            name
            for name in PRESERVED_PARTS
            if c14n(package[name]) != c14n(reference_package[name])
        ]
        output_sect = document_root.xpath("/w:document/w:body/w:sectPr", namespaces=NS)
        reference_sect = reference_document_root.xpath("/w:document/w:body/w:sectPr", namespaces=NS)
        section_equal = (
            len(output_sect) == len(reference_sect) == 1
            and c14n(output_sect[0]) == c14n(reference_sect[0])
        )
        metrics = page_metrics(document)
        record(
            "API",
            f"{variant} styles relationships section and A4",
            not preserved_mismatches and section_equal and page_metrics_valid(metrics),
            f"part_mismatches={preserved_mismatches}, section_equal={section_equal}, metrics={json.dumps(metrics, ensure_ascii=False)}",
        )

        header_root = parsed_xml["word/header1.xml"]
        reference_header_root = reference_xml["word/header1.xml"]
        header_text = "".join(header_root.xpath(".//w:t/text()", namespaces=NS))
        page_fields = [
            text.strip()
            for text in header_root.xpath(".//w:instrText/text()", namespaces=NS)
            if "PAGE" in text
        ]
        drawings = len(header_root.xpath(".//w:drawing", namespaces=NS))
        picts = len(header_root.xpath(".//w:pict", namespaces=NS))
        header_copy = etree.fromstring(package["word/header1.xml"])
        for text_node in header_copy.xpath(".//w:t", namespaces=NS):
            if text_node.text:
                text_node.text = text_node.text.replace("云技", "来绘")
        header_equal_after_brand = c14n(header_copy) == c14n(reference_header_root)
        old_brand_count = b"\n".join(package[name] for name in parsed_xml).count("来绘".encode("utf-8"))
        header_shape_ok = (
            len(page_fields) == spec["page_fields"]
            and drawings == spec["drawings"]
            and picts == spec["picts"]
        )
        if variant == "B":
            anchors = header_root.xpath(".//wp:anchor", namespaces=NS)
            header_shape_ok = header_shape_ok and len(anchors) == 1 and anchors[0].get("behindDoc") == "1"
        record(
            "API",
            f"{variant} header fidelity and PAGE fields",
            header_equal_after_brand and header_shape_ok,
            f"header={header_text!r}, PAGE={len(page_fields)}, drawings={drawings}, picts={picts}, equal_after_brand={header_equal_after_brand}",
        )
        record(
            "E2E",
            f"{variant} cloud-tech brand without old subject",
            "云技" in header_text and old_brand_count == 0,
            f"header={header_text!r}, old_brand_xml={old_brand_count}",
        )

        media_names = sorted(name for name in package if name.startswith("word/media/"))
        image = Image.open(io.BytesIO(package["word/media/image1.png"]))
        image.load()
        extrema = image.convert("RGB").getextrema()
        image_hash = sha256_bytes(package["word/media/image1.png"])
        record(
            "API",
            f"{variant} one white background",
            media_names == ["word/media/image1.png"]
            and image.size == (640, 960)
            and image.mode == "RGB"
            and extrema == ((255, 255), (255, 255), (255, 255)),
            f"media={media_names}, size={image.size}, mode={image.mode}, extrema={extrema}, sha256={image_hash}",
        )

        paragraphs = document.paragraphs
        boundaries_ok = (
            paragraphs[0].text == "程序开始！！！"
            and paragraphs[-1].text == "程序结束！！！"
        )
        content_mismatches: list[str] = []
        extracted_empty_lines = 0
        extracted_tabs = 0

        if variant == "A":
            position = 1
            for item in items:
                metadata = (
                    f"工程：{item.project}",
                    f"源文件：{item.relative_path}",
                    f"语言：{item.language}",
                    "源代码：",
                )
                if tuple(paragraph.text for paragraph in paragraphs[position : position + 4]) != metadata:
                    content_mismatches.append(item.relative_path + ":metadata")
                position += 4
                expected_lines = item.source_text.splitlines()
                actual_lines = [paragraph.text for paragraph in paragraphs[position : position + len(expected_lines)]]
                if actual_lines != expected_lines:
                    content_mismatches.append(item.relative_path + ":lines")
                extracted_empty_lines += sum(line == "" for line in actual_lines)
                extracted_tabs += sum(line.count("\t") for line in actual_lines)
                position += len(expected_lines)
                expected_end = f"文件结束：{item.relative_path}"
                if position >= len(paragraphs) or paragraphs[position].text != expected_end:
                    content_mismatches.append(item.relative_path + ":end")
                position += 1
            boundaries_ok = boundaries_ok and position == len(paragraphs) - 1

            body_paragraphs = document_root.xpath("/w:document/w:body/w:p", namespaces=NS)
            reference_body_paragraphs = reference_document_root.xpath(
                "/w:document/w:body/w:p", namespaces=NS
            )
            prototype_ppr = c14n(reference_body_paragraphs[1].find(f"{{{W_NS}}}pPr"))
            prototype_rpr = first_run_properties(reference_body_paragraphs[1])
            intermediate_ppr_ok = all(
                c14n(paragraph.find(f"{{{W_NS}}}pPr")) == prototype_ppr
                for paragraph in body_paragraphs[1:-1]
            )
            nonempty_run_rpr_ok = all(
                first_run_properties(paragraph) == prototype_rpr
                for paragraph in body_paragraphs[1:-1]
                if "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
            )
            edge_ppr_ok = (
                c14n(body_paragraphs[0].find(f"{{{W_NS}}}pPr"))
                == c14n(reference_body_paragraphs[0].find(f"{{{W_NS}}}pPr"))
                and c14n(body_paragraphs[-1].find(f"{{{W_NS}}}pPr"))
                == c14n(reference_body_paragraphs[-1].find(f"{{{W_NS}}}pPr"))
            )
            record(
                "API",
                "A line-paragraph style prototype",
                intermediate_ppr_ok and nonempty_run_rpr_ok and edge_ppr_ok,
                f"pPr={intermediate_ppr_ok}, rPr={nonempty_run_rpr_ok}, edges={edge_ppr_ok}, prototype_pPr_sha256={sha256_bytes(prototype_ppr)}, prototype_rPr_sha256={sha256_bytes(prototype_rpr)}",
            )
            break_count = len(document_root.xpath("//w:br", namespaces=NS))
            tab_count = len(document_root.xpath("//w:tab", namespaces=NS))
            record(
                "DATA",
                "A empty-line and tab preservation",
                extracted_empty_lines == source_empty_lines == 296
                and extracted_tabs == source_tabs == tab_count == 0
                and break_count == spec["breaks"],
                f"empty={extracted_empty_lines}/{source_empty_lines}, tabs={extracted_tabs}/{source_tabs}/{tab_count}, breaks={break_count}",
            )
        else:
            if len(paragraphs[1:-1]) != len(items):
                content_mismatches.append("container-count")
            for item, paragraph in zip(items, paragraphs[1:-1]):
                if paragraph.text != item.container_text:
                    content_mismatches.append(item.relative_path)
            ppr_equal = paragraph_properties(document_root) == paragraph_properties(reference_document_root)
            break_count = len(document_root.xpath("//w:br", namespaces=NS))
            record(
                "API",
                "B per-file paragraph property fidelity",
                ppr_equal and break_count == spec["breaks"],
                f"pPr_equal={ppr_equal}, textWrapping_breaks={break_count}",
            )

        record(
            "E2E",
            f"{variant} organization and boundaries",
            boundaries_ok and not content_mismatches,
            f"paragraphs={len(paragraphs)}, mismatches={content_mismatches}",
        )
        record(
            "DATA",
            f"{variant} source reconstruction",
            not content_mismatches,
            f"files={len(items)}, lines={source_lines}, chars={source_chars}, mismatches={content_mismatches}",
        )

        full_text = "\n".join(paragraph.text for paragraph in paragraphs)
        sensitive_hits = scan_credentials(full_text)
        replacement_count = full_text.count("\ufffd")
        control_chars = [
            ord(char)
            for char in full_text
            if ord(char) < 32 and char not in {"\n", "\r", "\t"}
        ]
        record("DATA", f"{variant} sensitive scan", not sensitive_hits, repr(sensitive_hits))
        record(
            "QUALITY",
            f"{variant} readable untruncated text",
            replacement_count == 0 and not control_chars and not content_mismatches,
            f"replacement={replacement_count}, controls={control_chars[:10]}, mismatches={content_mismatches}",
        )

        generator_source = generator_path.read_text(encoding="utf-8")
        generator_tree = ast.parse(generator_source, filename=str(generator_path))
        modules = imported_modules(generator_tree)
        calls = called_names(generator_tree)
        required_modules = {"docx", "lxml", "PIL", "zipfile"}
        required_calls = {"Document", "ZipFile", "fromstring"}
        record(
            "QUALITY",
            f"{variant} structured generator static review",
            required_modules.issubset(modules) and required_calls.issubset(calls),
            f"modules={sorted(required_modules & modules)}, calls={sorted(required_calls & calls)}",
        )
        record(
            "QUALITY",
            f"{variant} inventory and generator identity",
            actual_inventory_hash == spec["inventory_sha256"]
            and actual_generator_hash == spec["generator_sha256"],
            f"inventory={actual_inventory_hash}, generator={actual_generator_hash}",
        )
        expected_record_values = (
            str(spec["output_sha256"]),
            f"{int(spec['output_bytes']):,}",
            str(spec["reference_sha256"]),
            str(spec["inventory_sha256"]),
            str(spec["generator_sha256"]),
        )
        record(
            "QUALITY",
            f"{variant} generation record consistency",
            all(value in generation_record_text for value in expected_record_values),
            f"record={generation_record_path.name}, values_present={[(value in generation_record_text) for value in expected_record_values]}",
        )

        state[variant] = {
            "output_path": str(output_path),
            "output_sha256": actual_output_hash,
            "output_bytes": actual_output_bytes,
            "reference_sha256": actual_reference_hash,
            "inventory_sha256": actual_inventory_hash,
            "generator_sha256": actual_generator_hash,
            "paragraphs": len(paragraphs),
            "tables": len(document.tables),
            "sections": len(document.sections),
            "package_files": len(package),
            "xml_parts": len(parsed_xml),
            "source_files": len(items),
            "source_lines": source_lines,
            "source_chars": source_chars,
            "empty_source_lines": source_empty_lines,
            "source_tabs": source_tabs,
            "header_text": header_text,
            "page_fields": len(page_fields),
            "image_sha256": image_hash,
            "sensitive_hits": sensitive_hits,
            "content_mismatches": content_mismatches,
        }

    record(
        "DATA",
        "paired deliverables cover five projects",
        combined_projects == EXPECTED_PROJECTS,
        ", ".join(sorted(combined_projects)),
    )
    record(
        "E2E",
        "A/B one-to-one template organization",
        state["A"]["paragraphs"] == 2646
        and state["A"]["page_fields"] == 2
        and state["B"]["paragraphs"] == 20
        and state["B"]["page_fields"] == 0,
        f"A={state['A']['paragraphs']} paragraphs/{state['A']['page_fields']} PAGE, B={state['B']['paragraphs']} paragraphs/{state['B']['page_fields']} PAGE",
    )
    record(
        "QUALITY",
        "existing B artifact remained unchanged",
        state["B"]["output_sha256"] == EXPECTED["B"]["output_sha256"],
        str(state["B"]["output_sha256"]),
    )

    git_status = run_git(repo_root, "status", "--short", "--", "code/develop")
    git_diff = run_git(repo_root, "diff", "--", "code/develop")
    record(
        "QUALITY",
        "business source git status clean",
        git_status.returncode == 0 and not git_status.stdout.strip(),
        git_status.stdout.strip() or git_status.stderr.strip() or "empty",
    )
    record(
        "QUALITY",
        "business source git diff empty",
        git_diff.returncode == 0 and not git_diff.stdout,
        f"stdout_bytes={len(git_diff.stdout.encode('utf-8'))}, stderr={git_diff.stderr.strip()!r}",
    )

    groups = {
        group: all(check.passed for check in checks if check.group == group)
        for group in ("E2E", "DATA", "API", "QUALITY")
    }
    summary = {
        "checked_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "groups": groups,
        "artifacts": state,
        "paired_projects": sorted(combined_projects),
        "git": {
            "status_code_develop_empty": not git_status.stdout.strip(),
            "diff_code_develop_empty": not git_diff.stdout,
        },
        "office_renderers": detect_office_renderers(),
        "checks": {"passed": sum(check.passed for check in checks), "total": len(checks)},
    }

    for check in checks:
        status = "PASS" if check.passed else "FAIL"
        print(f"[{status}] {check.group}.{check.name}: {check.detail}")
    print("SUMMARY_JSON=" + json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0 if all(groups.values()) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"[ERROR] verifier exception: {type(exc).__name__}: {exc}", file=sys.stderr)
        raise
