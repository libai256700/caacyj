from __future__ import annotations

import ast
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image


EXPECTED_DOCX_NAME = "云技_软件设计说明书源代码_1.docx"
EXPECTED_DOCX_SHA256 = "0530f54d3f1936daaaabb62a59c07fca6201461d9b421904f80911e79a27dcb1"
EXPECTED_DOCX_SIZE = 80660
EXPECTED_TEMPLATE_SHA256 = "3271e0e1a64ad834808723230995948df5b23f9988adf8c0dd96a4b2f16301b2"
EXPECTED_PROJECTS = {
    "yunjikeji",
    "yunjikeji-admin-ui",
    "yunjikeji-server",
    "yunjikeji-admin-server",
    "fly-llm",
}
EXPECTED_SOURCE_FILES = 18
EXPECTED_SOURCE_LINES = 8329
EXPECTED_SOURCE_CHARS = 310972
EXPECTED_PARAGRAPHS = 20
EXPECTED_PACKAGE_FILES = 15
EXPECTED_XML_PARTS = 14

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "wp": WP_NS, "pr": REL_NS}

REQUIRED_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/_rels/document.xml.rels",
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
    "word/settings.xml",
    "word/header1.xml",
    "word/_rels/header1.xml.rels",
    "word/media/image1.png",
}
PRESERVED_XML_PARTS = (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
)
EXCLUDED_PATH_MARKERS = (
    "/.env",
    "/tmp/",
    "/node_modules/",
    "/dist/",
    "/target/",
    "/unpackage/",
    "application.yml",
    "application.yaml",
    "application-",
    "账号",
    "密码",
)
EXCLUDED_SUFFIXES = {
    ".pem",
    ".key",
    ".p12",
    ".pfx",
    ".jks",
    ".crt",
    ".cer",
    ".sql",
}
CREDENTIAL_PATTERNS = (
    ("private-key", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    ("certificate", re.compile(r"-----BEGIN CERTIFICATE-----")),
    ("aws-access-key", re.compile(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b", re.IGNORECASE)),
    ("long-sk-key", re.compile(r"\bsk-[A-Za-z0-9_-]{16,}\b", re.IGNORECASE)),
    (
        "literal-sensitive-assignment",
        re.compile(
            r"\b(?:password|passwd|pwd|secret|token|api[_-]?key|apikey|"
            r"access[_-]?key|secret[_-]?key|private[_-]?key)\b\s*[:=]\s*"
            r"[\"'][^\"'\r\n]{4,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "literal-authorization",
        re.compile(
            r"\bauthorization\b\s*[:=]\s*[\"'](?:bearer|basic)\s+"
            r"[A-Za-z0-9._+/=-]{8,}[\"']",
            re.IGNORECASE,
        ),
    ),
    (
        "credential-in-url",
        re.compile(r"\b[a-z][a-z0-9+.-]*://[^\s/:]+:[^\s/@]+@", re.IGNORECASE),
    ),
)


@dataclass(frozen=True)
class SourceItem:
    order: int
    project: str
    relative_path: str
    language: str
    expected_lines: int
    expected_chars: int
    expected_sha256: str
    source_text: str

    @property
    def container_text(self) -> str:
        return (
            f"工程：{self.project}\n"
            f"路径：{self.relative_path}\n"
            f"语言：{self.language}\n"
            "源码：\n"
            + normalize_newlines(self.source_text)
        )


@dataclass(frozen=True)
class Check:
    group: str
    name: str
    passed: bool
    detail: str


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def normalize_newlines(text: str) -> str:
    return re.sub(r"\r\n|\r|\n", "\n", text)


def utf16_character_count(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def strip_ticks(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == "`" and value[-1] == "`":
        return value[1:-1]
    return value


def find_repo_root(start: Path) -> Path:
    for candidate in (start.parent, *start.parents):
        if (candidate / ".git").exists() and (candidate / "code" / "develop").is_dir():
            return candidate.resolve()
    raise RuntimeError("Cannot locate repository root")


def is_within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def load_inventory(path: Path, repo_root: Path) -> list[SourceItem]:
    items: list[SourceItem] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 9 or not cells[0].isdigit():
            continue
        relative_path = strip_ticks(cells[2]).replace("\\", "/")
        source_path = (repo_root / relative_path).resolve()
        if not relative_path.startswith("code/develop/") or not is_within(source_path, repo_root):
            raise RuntimeError(f"Inventory path outside repository: {relative_path}")
        raw = source_path.read_bytes()
        source_text = raw.decode("utf-8-sig", errors="strict")
        items.append(
            SourceItem(
                order=int(cells[0]),
                project=strip_ticks(cells[1]),
                relative_path=relative_path,
                language=strip_ticks(cells[3]),
                expected_lines=int(cells[5].replace(",", "")),
                expected_chars=int(cells[6].replace(",", "")),
                expected_sha256=strip_ticks(cells[7]).lower(),
                source_text=source_text,
            )
        )
    return items


def c14n(data_or_element: bytes | etree._Element | None) -> bytes:
    if data_or_element is None:
        return b"<none>"
    element = (
        etree.fromstring(data_or_element)
        if isinstance(data_or_element, bytes)
        else data_or_element
    )
    return etree.tostring(element, method="c14n", with_comments=False)


def xml_parts(package: dict[str, bytes]) -> dict[str, etree._Element]:
    parsed: dict[str, etree._Element] = {}
    for name, data in package.items():
        if name.endswith((".xml", ".rels")):
            parsed[name] = etree.fromstring(data)
    return parsed


def paragraph_properties(document_root: etree._Element) -> list[bytes]:
    paragraphs = document_root.xpath("/w:document/w:body/w:p", namespaces=NS)
    return [c14n(paragraph.find(f"{{{W_NS}}}pPr")) for paragraph in paragraphs]


def run_git(repo_root: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        ["git", *args],
        cwd=repo_root,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )


def imported_modules(tree: ast.AST) -> set[str]:
    modules: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            modules.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            modules.add(node.module.split(".")[0])
    return modules


def called_names(tree: ast.AST) -> set[str]:
    names: set[str] = set()
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        if isinstance(node.func, ast.Name):
            names.add(node.func.id)
        elif isinstance(node.func, ast.Attribute):
            names.add(node.func.attr)
    return names


def detect_office_renderers() -> list[str]:
    found: list[str] = []
    for executable in ("WINWORD.EXE", "wps.exe", "soffice.exe", "libreoffice.exe"):
        path = shutil.which(executable)
        if path:
            found.append(path)
    for path in (
        Path(r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE"),
        Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE"),
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ):
        if path.is_file():
            found.append(str(path))
    return sorted(set(found))


def main() -> int:
    script_path = Path(__file__).resolve()
    output_dir = script_path.parent
    repo_root = find_repo_root(script_path)
    inventory_path = output_dir / "02-源码选取清单.md"
    generation_record_path = output_dir / "03-生成记录.md"
    generator_path = output_dir / "generate_source_design_docx.py"
    target_path = Path.home() / "Desktop" / EXPECTED_DOCX_NAME
    template_path = (
        Path.home()
        / "Downloads"
        / "飞书20260724-105710"
        / "来绘_软件设计说明书源代码_1.docx"
    )

    checks: list[Check] = []

    def record(group: str, name: str, passed: bool, detail: str) -> None:
        checks.append(Check(group, name, bool(passed), detail))

    actual_docx_hash = sha256_file(target_path) if target_path.is_file() else "missing"
    actual_docx_size = target_path.stat().st_size if target_path.is_file() else -1
    deliverables = sorted(p.name for p in target_path.parent.glob("云技_软件设计说明书源代码*.docx"))
    record("E2E", "exact deliverable", target_path.is_file(), str(target_path))
    record("E2E", "single requested filename", deliverables == [EXPECTED_DOCX_NAME], repr(deliverables))
    record(
        "E2E",
        "independent hash and size",
        actual_docx_hash == EXPECTED_DOCX_SHA256 and actual_docx_size == EXPECTED_DOCX_SIZE,
        f"sha256={actual_docx_hash}, bytes={actual_docx_size}",
    )

    inventory_items = load_inventory(inventory_path, repo_root)
    inventory_hash = sha256_file(inventory_path)
    orders = [item.order for item in inventory_items]
    projects = {item.project for item in inventory_items}
    record(
        "DATA",
        "inventory cardinality and order",
        len(inventory_items) == EXPECTED_SOURCE_FILES and orders == list(range(1, 19)),
        f"files={len(inventory_items)}, orders={orders}",
    )
    record(
        "DATA",
        "five-project coverage",
        projects == EXPECTED_PROJECTS,
        ", ".join(sorted(projects)),
    )

    missing_paths: list[str] = []
    hash_mismatches: list[str] = []
    stat_mismatches: list[str] = []
    excluded_paths: list[str] = []
    total_lines = 0
    total_chars = 0
    for item in inventory_items:
        source_path = repo_root / item.relative_path
        if not source_path.is_file():
            missing_paths.append(item.relative_path)
            continue
        raw = source_path.read_bytes()
        if sha256_bytes(raw) != item.expected_sha256:
            hash_mismatches.append(item.relative_path)
        actual_lines = len(item.source_text.splitlines())
        actual_chars = utf16_character_count(item.source_text)
        total_lines += actual_lines
        total_chars += actual_chars
        if (actual_lines, actual_chars) != (item.expected_lines, item.expected_chars):
            stat_mismatches.append(
                f"{item.relative_path}: {actual_lines}/{actual_chars} != "
                f"{item.expected_lines}/{item.expected_chars}"
            )
        lowered = "/" + item.relative_path.lower()
        if any(marker in lowered for marker in EXCLUDED_PATH_MARKERS) or source_path.suffix.lower() in EXCLUDED_SUFFIXES:
            excluded_paths.append(item.relative_path)

    record("DATA", "all source paths exist", not missing_paths, repr(missing_paths))
    record("DATA", "source hashes match inventory", not hash_mismatches, repr(hash_mismatches))
    record("DATA", "source line/character metadata", not stat_mismatches, repr(stat_mismatches))
    record(
        "DATA",
        "source aggregate",
        (total_lines, total_chars) == (EXPECTED_SOURCE_LINES, EXPECTED_SOURCE_CHARS),
        f"lines={total_lines}, chars={total_chars}",
    )
    record("DATA", "excluded-path policy", not excluded_paths, repr(excluded_paths))

    with zipfile.ZipFile(target_path) as archive:
        crc_error = archive.testzip()
        package = {name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}
    parsed_xml = xml_parts(package)
    record("API", "ZIP CRC", crc_error is None, repr(crc_error))
    record(
        "API",
        "package parts",
        len(package) == EXPECTED_PACKAGE_FILES and REQUIRED_PARTS.issubset(package),
        f"files={len(package)}, missing={sorted(REQUIRED_PARTS - package.keys())}",
    )
    record(
        "API",
        "all XML and relationships parse",
        len(parsed_xml) == EXPECTED_XML_PARTS,
        f"parsed={len(parsed_xml)}",
    )

    document = Document(target_path)
    paragraphs = document.paragraphs
    record("E2E", "python-docx reopen", True, "opened successfully")
    record(
        "E2E",
        "reference paragraph organization",
        len(paragraphs) == EXPECTED_PARAGRAPHS
        and len(document.tables) == 0
        and len(document.sections) == 1,
        f"paragraphs={len(paragraphs)}, tables={len(document.tables)}, sections={len(document.sections)}",
    )
    record(
        "E2E",
        "program boundaries",
        paragraphs[0].text == "程序开始！！！" and paragraphs[-1].text == "程序结束！！！",
        f"first={paragraphs[0].text!r}, last={paragraphs[-1].text!r}",
    )

    content_mismatches: list[str] = []
    for item, paragraph in zip(inventory_items, paragraphs[1:-1], strict=True):
        if paragraph.text != item.container_text:
            content_mismatches.append(item.relative_path)
    record(
        "DATA",
        "18 metadata/source containers exact",
        len(paragraphs[1:-1]) == EXPECTED_SOURCE_FILES and not content_mismatches,
        f"containers={len(paragraphs[1:-1])}, mismatches={content_mismatches}",
    )

    full_text = "\n".join(paragraph.text for paragraph in paragraphs)
    sensitive_hits = [name for name, pattern in CREDENTIAL_PATTERNS if pattern.search(full_text)]
    record("DATA", "credential and certificate scan", not sensitive_hits, repr(sensitive_hits))

    xml_payload = b"\n".join(package[name] for name in sorted(parsed_xml))
    old_brand_count = xml_payload.count("来绘".encode("utf-8"))
    header_root = parsed_xml["word/header1.xml"]
    header_text = "".join(header_root.xpath(".//w:t/text()", namespaces=NS))
    drawing_count = len(header_root.xpath(".//w:drawing", namespaces=NS))
    anchors = header_root.xpath(".//wp:anchor", namespaces=NS)
    behind_doc = anchors[0].get("behindDoc") if len(anchors) == 1 else None
    record(
        "E2E",
        "cloud-tech brand and no old subject",
        "云技" in header_text and old_brand_count == 0 and projects == EXPECTED_PROJECTS,
        f"header={header_text!r}, old_brand_xml={old_brand_count}",
    )
    record(
        "API",
        "header drawing preserved",
        drawing_count == 1 and len(anchors) == 1 and behind_doc == "1",
        f"drawings={drawing_count}, anchors={len(anchors)}, behindDoc={behind_doc}",
    )

    media_names = sorted(name for name in package if name.startswith("word/media/"))
    image = Image.open(io.BytesIO(package["word/media/image1.png"]))
    image.load()
    rgb = image.convert("RGB")
    extrema = rgb.getextrema()
    image_hash = sha256_bytes(package["word/media/image1.png"])
    record(
        "API",
        "one white 640x960 RGB background",
        media_names == ["word/media/image1.png"]
        and image.size == (640, 960)
        and image.mode == "RGB"
        and extrema == ((255, 255), (255, 255), (255, 255)),
        f"media={media_names}, size={image.size}, mode={image.mode}, extrema={extrema}, sha256={image_hash}",
    )

    template_hash = sha256_file(template_path)
    with zipfile.ZipFile(template_path) as archive:
        template_package = {
            name: archive.read(name) for name in archive.namelist() if not name.endswith("/")
        }
    template_xml = xml_parts(template_package)
    preserved_mismatches = [
        name
        for name in PRESERVED_XML_PARTS
        if c14n(package[name]) != c14n(template_package[name])
    ]
    output_doc_root = parsed_xml["word/document.xml"]
    template_doc_root = template_xml["word/document.xml"]
    output_sect = output_doc_root.xpath("/w:document/w:body/w:sectPr", namespaces=NS)
    template_sect = template_doc_root.xpath("/w:document/w:body/w:sectPr", namespaces=NS)
    section_equal = (
        len(output_sect) == len(template_sect) == 1 and c14n(output_sect[0]) == c14n(template_sect[0])
    )
    ppr_equal = paragraph_properties(output_doc_root) == paragraph_properties(template_doc_root)
    table_nodes = output_doc_root.xpath("//w:tbl", namespaces=NS)
    record(
        "API",
        "template style/theme/font fidelity",
        template_hash == EXPECTED_TEMPLATE_SHA256 and not preserved_mismatches,
        f"template_sha256={template_hash}, mismatches={preserved_mismatches}",
    )
    record(
        "API",
        "section and paragraph-property fidelity",
        section_equal and ppr_equal and not table_nodes,
        f"section_equal={section_equal}, pPr_equal={ppr_equal}, tables={len(table_nodes)}",
    )

    section = document.sections[0]
    to_mm = lambda value: round(float(value) / 36000.0, 3)
    page_metrics = {
        "width_mm": to_mm(section.page_width),
        "height_mm": to_mm(section.page_height),
        "top_mm": to_mm(section.top_margin),
        "bottom_mm": to_mm(section.bottom_margin),
        "left_mm": to_mm(section.left_margin),
        "right_mm": to_mm(section.right_margin),
        "header_mm": to_mm(section.header_distance),
        "footer_mm": to_mm(section.footer_distance),
    }
    expected_metrics = {
        "width_mm": 210.0,
        "height_mm": 297.0,
        "top_mm": 10.0,
        "bottom_mm": 10.0,
        "left_mm": 31.75,
        "right_mm": 31.75,
        "header_mm": 12.7,
        "footer_mm": 12.7,
    }
    metrics_ok = all(abs(page_metrics[key] - value) <= 0.2 for key, value in expected_metrics.items())
    record("API", "A4 and margin metrics", metrics_ok, json.dumps(page_metrics, ensure_ascii=False))

    record_text = generation_record_path.read_text(encoding="utf-8")
    record_size_match = re.search(r"成品大小：`([\d,]+)` 字节", record_text)
    record_hash_match = re.search(r"成品 SHA-256：`([0-9a-f]{64})`", record_text)
    record_size = int(record_size_match.group(1).replace(",", "")) if record_size_match else -1
    record_hash = record_hash_match.group(1) if record_hash_match else "missing"
    manifest_asset_line = next(
        (line for line in record_text.splitlines() if line.startswith("| 源码清单 |")),
        "",
    )
    generator_asset_line = next(
        (line for line in record_text.splitlines() if line.startswith("| 生成脚本 |")),
        "",
    )
    manifest_record_hashes = re.findall(r"`([0-9a-f]{64})`", manifest_asset_line)
    generator_record_hashes = re.findall(r"`([0-9a-f]{64})`", generator_asset_line)
    generator_hash = sha256_file(generator_path)
    record(
        "QUALITY",
        "generation record matches artifact",
        record_size == actual_docx_size and record_hash == actual_docx_hash,
        f"record={record_hash}/{record_size}, actual={actual_docx_hash}/{actual_docx_size}",
    )
    record(
        "QUALITY",
        "recorded manifest/generator hashes",
        manifest_record_hashes == [inventory_hash]
        and generator_record_hashes == [generator_hash],
        f"manifest={inventory_hash}, generator={generator_hash}",
    )

    generator_source = generator_path.read_text(encoding="utf-8")
    generator_tree = ast.parse(generator_source, filename=str(generator_path))
    modules = imported_modules(generator_tree)
    calls = called_names(generator_tree)
    structured_modules = {"docx", "lxml", "PIL", "zipfile"}
    structured_calls = {"Document", "ZipFile", "fromstring"}
    record(
        "QUALITY",
        "structured generator static inspection",
        structured_modules.issubset(modules) and structured_calls.issubset(calls),
        f"modules={sorted(structured_modules & modules)}, calls={sorted(structured_calls & calls)}",
    )

    git_status = run_git(repo_root, "status", "--short", "--", "code/develop")
    git_diff = run_git(repo_root, "diff", "--", "code/develop")
    record(
        "QUALITY",
        "business source git status clean",
        git_status.returncode == 0 and not git_status.stdout.strip(),
        git_status.stdout.strip() or git_status.stderr.strip() or "empty",
    )
    record(
        "QUALITY",
        "business source git diff empty",
        git_diff.returncode == 0 and not git_diff.stdout,
        f"stdout_bytes={len(git_diff.stdout.encode('utf-8'))}, stderr={git_diff.stderr.strip()!r}",
    )
    replacement_count = full_text.count("\ufffd")
    control_chars = [
        ord(char)
        for char in full_text
        if ord(char) < 32 and char not in {"\n", "\r", "\t"}
    ]
    record(
        "QUALITY",
        "readable text without replacement/control corruption",
        replacement_count == 0 and not control_chars and not content_mismatches,
        f"replacement_chars={replacement_count}, controls={control_chars[:10]}",
    )

    renderers = detect_office_renderers()
    groups = {
        group: all(check.passed for check in checks if check.group == group)
        for group in ("E2E", "DATA", "API", "QUALITY")
    }
    summary = {
        "checked_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "groups": groups,
        "docx": {
            "path": str(target_path),
            "sha256": actual_docx_hash,
            "bytes": actual_docx_size,
            "paragraphs": len(paragraphs),
            "tables": len(document.tables),
            "sections": len(document.sections),
            "package_files": len(package),
            "xml_parts": len(parsed_xml),
        },
        "sources": {
            "files": len(inventory_items),
            "projects": sorted(projects),
            "lines": total_lines,
            "chars": total_chars,
            "content_mismatches": content_mismatches,
        },
        "brand": {"header": header_text, "old_brand_xml": old_brand_count},
        "sensitive_hits": sensitive_hits,
        "git": {
            "status_code_develop_empty": not git_status.stdout.strip(),
            "diff_code_develop_empty": not git_diff.stdout,
        },
        "office_renderers": renderers,
        "checks": {"passed": sum(c.passed for c in checks), "total": len(checks)},
    }

    for check in checks:
        status = "PASS" if check.passed else "FAIL"
        print(f"[{status}] {check.group}.{check.name}: {check.detail}")
    print("SUMMARY_JSON=" + json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0 if all(groups.values()) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"[ERROR] verifier exception: {type(exc).__name__}: {exc}", file=sys.stderr)
        raise
